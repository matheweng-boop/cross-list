"""Excel (colour-highlighted) and CSV exports."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.matching.engine import CrossResult
from app.services.compare import Comparison

FILLS = {
    "better": PatternFill("solid", fgColor="C6EFCE"),
    "worse": PatternFill("solid", fgColor="FFC7CE"),
    "different": PatternFill("solid", fgColor="FFEB9C"),
    "unknown": PatternFill("solid", fgColor="E7E6E6"),
    "baseline": PatternFill("solid", fgColor="DDEBF7"),
}
BOLD = Font(bold=True)
SYNTHETIC_NOTE = "SYNTHETIC DEMO DATA – not real parts"


def _cross_rows(r: CrossResult) -> list[list[str]]:
    rows = []
    for d in r.all_diffs:
        rows.append([f"P{d.priority}", d.label, d.ref_display, d.cand_display, d.status])
    return rows


def _autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 60)


def cross_xlsx(r: CrossResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cross"
    ws.append([f"Cross: {r.reference.label} → {r.candidate.label}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Categories", ", ".join(r.categories) or "none"])
    if r.fff_status:
        ws.append(["FFF status", r.fff_status + (f" – {r.feedback.note}" if r.feedback and r.feedback.note else "")])
    ws.append(["Top 5 differences", r.comment()])
    if r.p3_notes():
        ws.append(
            ["Priority 3 notes", "; ".join(f"{d.label}: {d.cand_display} vs {d.ref_display}" for d in r.p3_notes())]
        )
    ws.append(["Candidate lifecycle", r.candidate.lifecycle_status.value])
    if r.reference.is_synthetic or r.candidate.is_synthetic:
        ws.append([SYNTHETIC_NOTE])
    ws.append([])
    header = ["Priority", "Parameter", r.reference.label + " (reference)", r.candidate.label, "Status"]
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = BOLD
    for row in _cross_rows(r):
        ws.append(row)
        status = row[-1]
        if status in FILLS:
            for c in ws[ws.max_row][2:]:
                c.fill = FILLS[status]
    ws["B3"].alignment = Alignment(wrap_text=True)
    _autosize(ws)
    ws.column_dimensions["B"].width = 30
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cross_csv(r: CrossResult) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Priority", "Parameter", r.reference.label + " (reference)", r.candidate.label, "Status"])
    w.writerows(_cross_rows(r))
    return buf.getvalue()


def _comparison_header(c: Comparison) -> list[str]:
    return ["Priority", "Parameter"] + [p.label + (" (baseline)" if p.id == c.baseline.id else "") for p in c.parts]


def _prio(param) -> str:
    return f"P{param.priority}" if param.priority else "DigiKey"


def comparison_xlsx(c: Comparison) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.append(["Comparison – priority 1 & 2 parameters"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Colours vs baseline: green better, red worse, yellow different, grey unknown. Bold = best in row."])
    if any(p.is_synthetic for p in c.parts):
        ws.append([SYNTHETIC_NOTE])
    ws.append([])
    ws.append(_comparison_header(c))
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    for row in c.rows:
        ws.append([_prio(row.param), row.param.label] + [cell.display for cell in row.cells])
        for xl, cell in zip(ws[ws.max_row][2:], row.cells, strict=True):
            if cell.status in FILLS:
                xl.fill = FILLS[cell.status]
            if cell.best:
                xl.font = BOLD
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def comparison_csv(c: Comparison) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(_comparison_header(c) + ["(best in row)"])
    for row in c.rows:
        best = [p.label for p, cell in zip(c.parts, row.cells, strict=True) if cell.best]
        values = [cell.display for cell in row.cells]
        w.writerow([_prio(row.param), row.param.label, *values, "; ".join(best)])
    return buf.getvalue()


RESULT_HEADER = ["Part", "Manufacturer", "Cross type", "Pin-to-pin", "Priority 1", "P1 better", "P1 worse or different",
                 "P2/P3 better (info)", "P2/P3 worse (info)", "Package", "Lifecycle", "Top 5 differences",
                 "Priority 3 notes"]  # fmt: skip
VERDICT_FILLS = {"fff": FILLS["better"], "functional": FILLS["baseline"], "package": FILLS["different"]}


def _result_cells(r: CrossResult, v) -> list[str]:
    return [
        r.candidate.mpn,
        r.candidate.manufacturer.name,
        v.label + (" (FFF rejected in review)" if r.fff_status == "rejected" else ""),
        "yes" if v.p2p else "no",
        v.spec_label,
        ", ".join(d.label for d in v.better),
        ", ".join(d.label for d in v.worse),
        ", ".join(d.label for d in v.info_better),
        ", ".join(d.label for d in v.info_worse),
        r.candidate.package.code,
        r.candidate.lifecycle_status.value,
        " · ".join(d.summary() for d in r.top_differences()),
        "; ".join(f"{d.label} {d.cand_display} vs {d.ref_display}" for d in r.p3_notes()),
    ]


def results_xlsx(reference, rows) -> bytes:
    """The part page's result table: one row per candidate with its drop-in verdict."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Replacements"
    ws.append([f"Replacements for {reference.mpn} ({reference.manufacturer.name}, {reference.package.code})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(
        [
            "Drop-in = same package, pin-to-pin, priority 1 equal or better. Functional = different package, "
            "few priority 1 deviations. Package eq. = same package, different spec. Priority 2/3 for information."
        ]
    )
    if reference.is_synthetic:
        ws.append([SYNTHETIC_NOTE])
    ws.append([])
    ws.append(RESULT_HEADER)
    for cell in ws[ws.max_row]:
        cell.font = BOLD
    for r, v in rows:
        ws.append(_result_cells(r, v))
        ws.cell(ws.max_row, 3).fill = VERDICT_FILLS[v.category]
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def results_csv(rows) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(RESULT_HEADER)
    for r, v in rows:
        w.writerow(_result_cells(r, v))
    return buf.getvalue()


# Products › Cross a list (owner, 26 Oct 2026): the on-screen table, light as text and fill colour, and a Summary sheet
LIGHT_FILLS = {"Green": PatternFill("solid", fgColor="C6EFCE"), "Amber": PatternFill("solid", fgColor="FFEB9C"),
               "Red": PatternFill("solid", fgColor="FFC7CE")}  # fmt: skip
CROSS_LIST_WIDTHS = {"#": 5, "Your part": 20, "Maker": 14, "Status": 11, "Light": 8, "Drop-in part": 20,
                     "Drop-in maker": 12, "Drop-in type": 22, "Drop-in comment": 60, "Functional fit part": 20,
                     "Functional fit maker": 12, "Functional fit type": 24, "Functional fit comment": 60,
                     "Checked on (P1 specs)": 40, "Verified": 11,
                     "Your comment": 30}  # fmt: skip


def cross_list_xlsx(rows, summary: dict, over: int = 0, accuracy: dict | None = None, shown: str = "") -> bytes:
    from app.services.crosslist import EXPORT_HEADER, MAX_LINES, PRECHECK_NOTE, export_lines, today

    wb = Workbook()
    ws = wb.active
    ws.title = "Cross list"
    ws.append(EXPORT_HEADER)
    for cell in ws[1]:
        cell.font = BOLD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    light_col = EXPORT_HEADER.index("Light") + 1
    wrap = {i + 1 for i, h in enumerate(EXPORT_HEADER) if "comment" in h.lower() or h.startswith("Checked")}
    for line in (x for r in rows for x in export_lines(r)):
        ws.append(line)
        n = ws.max_row
        ws.cell(n, light_col).fill = LIGHT_FILLS[ws.cell(n, light_col).value]
        for c in range(1, len(EXPORT_HEADER) + 1):
            ws.cell(n, c).alignment = Alignment(wrap_text=c in wrap, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADER))}{ws.max_row}"
    for i, h in enumerate(EXPORT_HEADER, 1):
        ws.column_dimensions[get_column_letter(i)].width = CROSS_LIST_WIDTHS.get(h, 14)

    sm = wb.create_sheet("Summary")
    lines = [["Cross a list", ""], ["Date", today()], ["Parts", summary["n"]], ["Green", summary["green"]],
             ["Amber", summary["amber"]], ["Red", summary["red"]], ["Coverage", f"{summary['coverage']} %"],
             ["Verified", summary["verified"]]]  # fmt: skip
    if shown:
        lines.append(["Rows exported", f"{shown} only"])
    if accuracy:
        lines.append(["Accuracy", f"Drop-ins correct: {accuracy.get('rate')} % of {accuracy.get('n')} known pairs "
                                  f"(checked {accuracy.get('checked')})"])  # fmt: skip
    if over:
        lines.append(["Left out", f"{over} lines beyond the first {MAX_LINES}"])
    if any(r.part is not None and r.part.is_synthetic for r in rows):
        lines.append(["Note", SYNTHETIC_NOTE])
    lines.append(["Note", PRECHECK_NOTE])
    for line in lines:
        sm.append(line)
    sm["A1"].font = Font(bold=True, size=13)
    sm.column_dimensions["A"].width, sm.column_dimensions["B"].width = 16, 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
