"""Presentation builder > Download as PowerPoint (owner, 25 Oct 2026): the same slides the preview shows
(app/services/presentation.build), written as an editable 16:9 .pptx with a small template defined here. One
PowerPoint slide per preview slide, in the same order and with the same header, title and footer. Text stays
text; parts tables are native PowerPoint tables; the system diagram (HTML blocks in the preview, no SVG) is drawn
as native boxes, so it can be edited too."""

import io
import re
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
W, H = Emu(12192000), Inches(7.5)  # 13.333 x 7.5 in, exactly 16:9
MX = Inches(0.6)  # side margin
INK, MUTED, ACCENT = RGBColor(0x1B, 0x1F, 0x27), RGBColor(0x5B, 0x62, 0x70), RGBColor(0x0B, 0x6B, 0x74)
SOFT, LINE, OURS, OURS_LINE = (RGBColor(0xE7, 0xF3, 0xF2), RGBColor(0xD7, 0xDB, 0xE3), RGBColor(0xDF, 0xF0, 0xEC),
                               RGBColor(0x7F, 0xB9, 0xAE))  # fmt: skip
SKY, QOR, WM = RGBColor(0x1F, 0x5F, 0xA8), RGBColor(0xC2, 0x41, 0x0C), RGBColor(0xEC, 0xEE, 0xF1)
FONT = "Calibri"
BODY_TOP, BODY_BOTTOM = Inches(1.55), Inches(6.85)


def filename(deck_title: str) -> str:
    """A safe ASCII file name from the deck title: "Skyworks + Qorvo" -> "skyworks-qorvo.pptx"."""
    s = re.sub(r"[^a-z0-9]+", "-", (deck_title or "").lower()).strip("-")[:60]
    return f"{s or 'presentation'}.pptx"


def _date(d) -> str:
    return f"{d.day} {d:%b %Y}" if isinstance(d, date) else str(d or "")


def _run(p, text: str, size: float, bold: bool = False, color: RGBColor = INK, mono: bool = False):
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    r.font.name = "Consolas" if mono else FONT
    return r


def _box(slide, x, y, w, h, *, fill: RGBColor | None = None, line: RGBColor | None = None, shape=MSO_SHAPE.RECTANGLE):
    if fill is None and line is None:
        sh = slide.shapes.add_textbox(x, y, w, h)
    else:
        sh = slide.shapes.add_shape(shape, x, y, w, h)
        sh.shadow.inherit = False
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb, sh.line.width = line, Pt(0.75)
    tf = sh.text_frame
    tf.word_wrap = True
    for side in ("left", "right"):
        setattr(tf, f"margin_{side}", Inches(0.08))
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return sh, tf


def _text(slide, x, y, w, h, text: str, size: float, *, bold=False, color=INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, **kw):  # fmt: skip
    sh, tf = _box(slide, x, y, w, h, **kw)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size, bold, color)
    return sh


def _fit(texts: list[str], base: float, chars: int) -> float:
    """A font size that keeps a block of texts inside the body: smaller as the text gets longer (est)."""
    n = sum(len(t) for t in texts) + 60 * len(texts)
    return max(base * 0.62, min(base, base * (chars / max(n, 1)) ** 0.5))


def _bullets(slide, x, y, w, h, points: list[str], size: float, *, lead: str = "", numbered_from: int = 0):
    _, tf = _box(slide, x, y, w, h)
    first = True
    if lead:
        p = tf.paragraphs[0]
        _run(p, lead, size * 1.05, color=MUTED)
        p.space_after = Pt(size * 0.6)
        first = False
    for i, pt in enumerate(points):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        mark = f"{numbered_from + i}.  " if numbered_from else "•  "
        _run(p, mark, size, True, ACCENT)
        _run(p, pt, size)
        p.space_after = Pt(size * 0.45)
    return tf


# ---- slide frame ----------------------------------------------------------------------------------------------------
def _frame(prs, deck_title: str, today, wm_text: str, n: int, total: int, logo: bytes | None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    if wm_text:  # first, so it sits behind everything
        wm = _text(slide, Inches(1), Inches(2.6), W - Inches(2), Inches(2.2), wm_text.upper(), 96, bold=True,
                   color=WM, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
        wm.rotation = -24
    _box(slide, 0, 0, W, Inches(0.55), fill=SOFT)
    _box(slide, 0, Inches(0.55), W, Inches(0.04), fill=ACCENT)
    x = MX
    if logo:
        try:
            pic = slide.shapes.add_picture(io.BytesIO(logo), MX, Inches(0.08), height=Inches(0.4))
            x = MX + pic.width + Inches(0.15)
        except Exception:  # a logo PowerPoint cannot read (SVG): leave it out
            pass
    _text(slide, x, Inches(0.06), W - x - MX, Inches(0.44), f"Skyworks + Qorvo · {deck_title}", 14, bold=True,
          color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    foot = f"{deck_title} · {_date(today)} · Public information" + (f" · {wm_text}" if wm_text else "")
    _box(slide, MX, Inches(6.98), W - 2 * MX, Emu(9525), fill=LINE)
    _text(slide, MX, Inches(7.02), W - 2 * MX - Inches(1.2), Inches(0.35), foot, 10, color=MUTED)
    _text(slide, W - MX - Inches(1.2), Inches(7.02), Inches(1.2), Inches(0.35), f"{n} / {total}", 10, color=MUTED,
          align=PP_ALIGN.RIGHT)  # fmt: skip
    return slide


def _title(slide, title: str):
    _text(slide, MX, Inches(0.75), W - 2 * MX, Inches(0.75), title, 28 if len(title) < 60 else 22, bold=True,
          anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip


# ---- slide bodies ---------------------------------------------------------------------------------------------------
def _cover(slide, b: dict):
    _text(slide, MX, Inches(2.3), W - 2 * MX, Inches(1.4), b.get("heading", ""), 48, bold=True, color=ACCENT,
          anchor=MSO_ANCHOR.BOTTOM)  # fmt: skip
    _text(slide, MX, Inches(3.8), W - 2 * MX, Inches(1.0), b.get("sub", ""), 24, color=MUTED)
    _text(slide, MX, Inches(4.9), W - 2 * MX, Inches(0.5), _date(b.get("date")), 16, color=MUTED)


def _points(slide, b: dict, kind: str):
    pts = list(b.get("points") or [])
    size = _fit([b.get("lead", ""), *pts], 22, 1400)
    _bullets(slide, MX, BODY_TOP, W - 2 * MX, BODY_BOTTOM - BODY_TOP, pts, size, lead=b.get("lead", ""),
             numbered_from=(b.get("start") or 1) if kind == "agenda" else 0)  # fmt: skip


def _diagram(slide, b: dict):
    _text(slide, MX, BODY_TOP - Inches(0.05), W - 2 * MX, Inches(0.4), b.get("name", ""), 14, color=MUTED)
    groups = b.get("groups") or []
    cols = max(1, min(int(b.get("cols") or len(groups) or 1), len(groups) or 1))
    rows = [groups[i : i + cols] for i in range(0, len(groups), cols)]
    gap = Inches(0.2)
    top, bottom = BODY_TOP + Inches(0.45), Inches(6.4)
    gw = (W - 2 * MX - gap * (cols - 1)) // cols
    rh = (bottom - top - gap * (len(rows) - 1)) // max(1, len(rows))
    most = max((len(g.get("blocks") or []) for g in groups), default=1)
    bh = min(Inches(0.55), (rh - Inches(0.5)) // max(1, most) - Inches(0.06))
    size = 12 if bh >= Inches(0.45) else 10
    for r, row in enumerate(rows):
        y = top + r * (rh + gap)
        for c, g in enumerate(row):
            x = MX + c * (gw + gap)
            _box(slide, x, y, gw, rh, fill=RGBColor(0xFA, 0xFB, 0xFD), line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            _text(slide, x + Inches(0.05), y + Inches(0.05), gw - Inches(0.1), Inches(0.4), g.get("label", ""), 13,
                  bold=True)  # fmt: skip
            by = y + Inches(0.5)
            for k in g.get("blocks") or []:
                ours = bool(k.get("ours"))
                _text(slide, x + Inches(0.1), by, gw - Inches(0.2), bh, k.get("label", ""), size, bold=ours,
                      fill=OURS if ours else RGBColor(0xFF, 0xFF, 0xFF), line=OURS_LINE if ours else LINE,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
                by += bh + Inches(0.06)
    ly = Inches(6.48)
    _text(slide, MX, ly, Inches(2.3), Inches(0.36), "Skyworks / Qorvo part", 11, bold=True, fill=OURS, line=OURS_LINE,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    _text(slide, MX + Inches(2.45), ly, Inches(1.6), Inches(0.36), "other blocks", 11, fill=RGBColor(0xFF, 0xFF, 0xFF),
          line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip


def _cell(cell, size: float, fill: RGBColor | None = None):
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    p = cell.text_frame.paragraphs[0]
    cell.text_frame.word_wrap = True
    return p


def _parts(slide, b: dict):
    rows = b.get("rows") or []
    auto = bool(b.get("automotive"))
    size = 14 if len(rows) <= 5 else 12
    height = min(BODY_BOTTOM - BODY_TOP - Inches(0.6), Inches(0.75) * (len(rows) + 1))
    gf = slide.shapes.add_table(len(rows) + 1, 3, MX, BODY_TOP, W - 2 * MX, height)
    t = gf.table
    t.columns[0].width = Inches(3.4)
    t.columns[1].width = t.columns[2].width = (W - 2 * MX - Inches(3.4)) // 2
    for i, head in enumerate(("Function", "Skyworks", "Qorvo")):
        _run(_cell(t.cell(0, i), size, ACCENT), head, size, True, RGBColor(0xFF, 0xFF, 0xFF))
    for r, row in enumerate(rows, start=1):
        bg = RGBColor(0xFF, 0xFF, 0xFF) if r % 2 else RGBColor(0xF6, 0xF8, 0xFA)
        _run(_cell(t.cell(r, 0), size, bg), row.get("fn", ""), size, True)
        for c, line in enumerate(("Skyworks", "Qorvo"), start=1):
            cell = t.cell(r, c)
            p = _cell(cell, size, bg)
            parts = (row.get("parts") or {}).get(line) or []
            for j, part in enumerate(parts):
                _run(p, part.get("mpn", ""), size, mono=True)
                if part.get("aec"):
                    _run(p, " AEC-Q", size * 0.75, True, RGBColor(0x3D, 0x61, 0x18))
                if j < len(parts) - 1:
                    _run(p, ", ", size)
            count = (row.get("counts") or {}).get(line)
            if count:
                q = cell.text_frame.add_paragraph()
                word = "part" if count == "1" else "parts"
                _run(q, f"{count} {'AEC-Q ' if auto else ''}{word} in the catalogue", size * 0.75, color=MUTED)
    note = "Representative parts" + ("; AEC-Q: automotive qualified" if auto else "") + \
        ". Counts: active catalogue parts, rounded."  # fmt: skip
    _text(slide, MX, BODY_BOTTOM - Inches(0.45), W - 2 * MX, Inches(0.4), note, 11, color=MUTED)


def _together(slide, b: dict):
    pairs = b.get("pairs") or []
    size = _fit([(p.get("line") or "") + p.get("skyworks", "") + p.get("qorvo", "") for p in pairs], 22, 1100)
    _, tf = _box(slide, MX, BODY_TOP, W - 2 * MX, BODY_BOTTOM - BODY_TOP)
    for i, pr in enumerate(pairs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _run(p, f"{i + 1}.  ", size, True, ACCENT)
        _run(p, "Skyworks ", size, True, SKY)
        _run(p, pr.get("skyworks", ""), size, True)
        _run(p, "  +  ", size, True, MUTED)
        _run(p, "Qorvo ", size, True, QOR)
        _run(p, pr.get("qorvo", ""), size, True)
        if pr.get("line"):
            q = tf.add_paragraph()
            _run(q, pr["line"], size * 0.9)
        if pr.get("s_ex") or pr.get("q_ex"):
            q = tf.add_paragraph()
            _run(q, "For example ", size * 0.75, color=MUTED)
            _run(q, pr.get("s_ex", ""), size * 0.75, mono=True)
            _run(q, " with ", size * 0.75, color=MUTED)
            _run(q, pr.get("q_ex", ""), size * 0.75, mono=True)
            if pr.get("n"):
                _run(q, f" · together in {pr['n']} reference designs", size * 0.75, color=MUTED)
        tf.paragraphs[-1].space_after = Pt(size * 0.8)


def _column(slide, x, y, w, h, head: str, items: list[str], size: float):
    _text(slide, x, y, w, Inches(0.45), head, size * 1.05, bold=True, color=ACCENT)
    _bullets(slide, x, y + Inches(0.45), w, h - Inches(0.45), items, size)


def _facts(slide, b: dict):
    cols = [(h, b.get(k) or []) for k, h in (("standards", "Standards"), ("bands", "Frequency bands")) if b.get(k)]
    specs = b.get("specs") or []
    size = _fit([*(x for _, v in cols for x in v), *specs], 18, 1300)
    top_h = (BODY_BOTTOM - BODY_TOP) * (0.5 if specs else 1)
    gap = Inches(0.4)
    cw = (W - 2 * MX - gap * (len(cols) - 1)) // max(1, len(cols))
    for i, (head, items) in enumerate(cols):  # two columns, as in the preview
        _column(slide, MX + i * (cw + gap), BODY_TOP, cw, top_h, head, items, size)
    if specs:
        y = BODY_TOP + (top_h if cols else 0)
        _column(slide, MX, y, W - 2 * MX, BODY_BOTTOM - y, "Typical specifications of our parts", specs, size)


def _glance(slide, b: dict):
    y = BODY_TOP
    if b.get("lead"):
        _text(slide, MX, y, W - 2 * MX, Inches(0.9), b["lead"], 16, color=MUTED)
        y += Inches(1.0)
    facts = b.get("facts") or []
    n = max(1, len(facts))
    gap = Inches(0.2)
    tw = (W - 2 * MX - gap * (n - 1)) // n
    for i, f in enumerate(facts):
        sh, tf = _box(slide, MX + i * (tw + gap), y, tw, Inches(1.6), fill=SOFT, line=None,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE)  # fmt: skip
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, f.get("num", ""), 32, True, ACCENT)
        q = tf.add_paragraph()
        q.alignment = PP_ALIGN.CENTER
        _run(q, f.get("label", ""), 14, color=INK)
    if b.get("pairs"):
        _, tf = _box(slide, MX, y + Inches(2.0), W - 2 * MX, Inches(1.6))
        p = tf.paragraphs[0]
        _run(p, "Where the two combine: ", 16, True)
        _run(p, " · ".join(b["pairs"]), 16)


def _closing(slide, b: dict):
    _text(slide, MX + Inches(0.6), Inches(2.3), W - 2 * MX - Inches(1.2), Inches(2.0), b.get("text", ""), 28,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    _text(slide, MX, Inches(4.6), W - 2 * MX, Inches(0.6), b.get("contact", ""), 18, color=ACCENT,
          align=PP_ALIGN.CENTER)  # fmt: skip


BODIES = {"bullets": _points, "agenda": _points, "diagram": _diagram, "parts": _parts, "together": _together,
          "facts": _facts, "glance": _glance, "closing": _closing}  # fmt: skip


def build(slides: list[dict], deck_title: str, today=None, wm_text: str = "", logo: bytes | None = None) -> bytes:
    """The .pptx bytes of the preview slides (presentation.build), one PowerPoint slide each."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prs.core_properties.title = f"Skyworks + Qorvo · {deck_title}"
    prs.core_properties.author = "Skyworks + Qorvo"
    prs.core_properties.comments = "Public information"
    today = today or date.today()
    total = len(slides)
    for n, sl in enumerate(slides, start=1):
        slide = _frame(prs, deck_title, today, wm_text, n, total, logo)
        kind, b = sl.get("kind"), sl.get("body") or {}
        if kind == "title":
            _cover(slide, b)
            continue
        _title(slide, sl.get("title", ""))
        if kind in ("bullets", "agenda"):
            _points(slide, b, kind)
        elif kind in BODIES:
            BODIES[kind](slide, b)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def texts(data: bytes) -> list[list[str]]:
    """Every text in a .pptx, slide by slide (shapes and table cells): for tests and the word check."""
    out = []
    for slide in Presentation(io.BytesIO(data)).slides:
        found = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                found.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                found += [c.text for row in sh.table.rows for c in row.cells]
        out.append(found)
    return out
