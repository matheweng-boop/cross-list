"""Browser flow tests (24 Oct 2026): the pages the owner uses, clicked in a real Chromium against the app on a real
port.

The other tests read the server's HTML; these catch what only happens in the browser (a script that does not fill a
field, a button that does nothing, a console error, a page wider than a phone). Every test also fails on any console
error or uncaught page error, and when a page it visited scrolls sideways at 390 px. Select them with
``-m browser``, leave them out with ``-m "not browser"``; they skip when the image has no Chromium.
"""

import re
from urllib.parse import quote

import pytest

from tests.conftest import PAGE_TIMEOUT

expect = pytest.importorskip("playwright.sync_api").expect

pytestmark = pytest.mark.browser

PHONE = {"width": 390, "height": 844}
WIDE = {"width": 1180, "height": 900}  # the iPad held sideways, a small Windows window


class Watch:
    """Console errors, uncaught page errors and the pages visited by one test."""

    def __init__(self, base: str):
        self.base, self.errors, self.visited = base, [], []

    def attach(self, page):
        page.on("console", lambda m: m.type == "error" and self.errors.append(f"console: {m.text} ({page.url})"))
        page.on("pageerror", lambda e: self.errors.append(f"page error: {e} ({page.url})"))
        page.on("framenavigated", lambda f: f == page.main_frame and self._visit(f.url))
        return page

    def _visit(self, url: str):
        url = url.split("#")[0]
        if url.startswith(self.base) and url not in self.visited:
            self.visited.append(url)


def overflow(page) -> int:
    """How many pixels the page is wider than the window (0: no sideways scrolling)."""
    return page.evaluate("Math.max(0, document.documentElement.scrollWidth - document.documentElement.clientWidth)")


@pytest.fixture
def page(chromium, live_server):
    """A fresh browser tab on the live app; afterwards every visited page is opened again at 390 px wide."""
    watch = Watch(live_server)
    ctx = chromium.new_context(base_url=live_server, viewport=WIDE)
    pg = watch.attach(ctx.new_page())
    pg.set_default_timeout(PAGE_TIMEOUT)
    yield pg
    ctx.close()
    phone = chromium.new_context(viewport=PHONE, is_mobile=True, has_touch=True)
    tab = watch.attach(phone.new_page())
    wide = []
    for url in list(watch.visited):
        tab.goto(url, wait_until="load")
        if (px := overflow(tab)) > 1:
            wide.append(f"{url} is {px} px wider than a 390 px phone")
    phone.close()
    problems = watch.errors + wide
    if problems:
        pytest.fail("\n".join(problems), pytrace=False)


def test_the_checks_themselves_catch_errors_and_overflow(chromium, live_server):
    """The watcher sees a console error and an uncaught error, and a page wider than a phone is measured as such."""
    watch = Watch(live_server)
    ctx = chromium.new_context(viewport=PHONE)
    tab = watch.attach(ctx.new_page())
    tab.goto(live_server + "/healthz")
    tab.evaluate("""() => { console.error("boom"); setTimeout(() => { throw new Error("bang"); });
                            const d = document.createElement("div"); d.style.width = "900px"; d.textContent = "x";
                            document.body.append(d); }""")
    tab.wait_for_timeout(100)
    assert overflow(tab) > 400
    ctx.close()
    assert any("boom" in e for e in watch.errors) and any("bang" in e for e in watch.errors)
    assert watch.visited == [live_server + "/healthz"]


def go(page, path: str):
    r = page.goto(path, wait_until="load")
    assert r is not None and r.status == 200, (path, r and r.status)
    return r


def test_every_menu_page_opens_without_errors(page):
    """Every page in the menus and tabs: no console error, and (checked after the test) none wider than a phone."""
    from tests.test_blank_forms import menu_pages

    for path in menu_pages():
        go(page, path)
        assert page.locator("main h1").count() >= 1, path


# --- Products ---------------------------------------------------------------------------------------------------------
def test_price_estimate_picks_a_part_override_and_estimate(page):
    go(page, "/pricing/estimate")
    form = page.locator("form.pe-form")
    form.locator("input[name=name]").fill("Si47962")
    form.locator("input[name=name]").press("Tab")  # leaving the box, as after picking from the list
    pins = form.locator("input[name=pins]")
    pins.wait_for()
    expect(pins).to_have_value("56")  # filled at once, before any submit
    assert form.locator("select[name=family]").input_value() == "am_fm_tuner"
    assert form.locator("input[name=package]").input_value() == "QFN-56"
    for name in ("family", "package", "pins"):  # locked: the part data fills them, "Override" is the second option
        assert "pe-locked" in (form.locator(f"[name={name}]").get_attribute("class") or ""), name
    assert form.locator("input[name=package]").evaluate("e => e.readOnly")
    assert form.locator("input[name=pins]").evaluate("e => e.readOnly")
    assert page.locator(".pe-from[data-for=pins]").inner_text() == "(from the part data)"

    page.get_by_role("button", name="Override").click()
    assert not form.locator("input[name=package]").evaluate("e => e.readOnly")
    assert "pe-locked" not in (form.locator("select[name=family]").get_attribute("class") or "")
    other = form.locator("select[name=family] option").nth(0)
    value, label = other.get_attribute("value"), other.inner_text().split(" · ")[0].strip()
    assert value != "am_fm_tuner"
    form.locator("select[name=family]").select_option(value)
    form.locator("input[name=bq]").first.fill("1000")
    form.locator("input[name=bp]").first.fill("2.64")
    with page.expect_navigation():
        form.get_by_role("button", name="Estimate").click()
    result = page.locator("#price-model.pe-result h2")
    result.wait_for()
    assert label in result.inner_text()  # the family I picked, not the part data's
    assert "override=true" in page.url and f"family={value}" in page.url
    assert page.locator("form.pe-form select[name=family]").input_value() == value


def test_price_estimate_with_every_number_blank(page):
    go(page, "/pricing/estimate")
    form = page.locator("form.pe-form")
    for name in ("pins", "sop", "released"):
        form.locator(f"input[name={name}]").fill("")
    with page.expect_navigation():
        form.get_by_role("button", name="Estimate").click()
    assert page.locator("main h1").inner_text() == "Price estimate"
    assert "released=" in page.url
    assert "cannot use" not in page.content()


def test_search_lists_matches_as_you_type_and_on_find(page):
    go(page, "/")
    box = page.locator("#finder-search")
    box.press_sequentially("SA77", delay=30)
    page.locator("#suggestions .suggest-list li").first.wait_for()
    assert "SA77" in page.locator("#suggestions").inner_text()
    with page.expect_navigation():
        box.press("Enter")
    assert "/find" in page.url
    assert page.locator("main").get_by_role("link", name=re.compile("SA77")).count() >= 2


def test_part_walk_through_next_step(page):
    go(page, "/find?q=SA7741")
    part_url = page.url.split("?")[0]
    assert re.search(r"/parts/\d+$", part_url)
    go(page, part_url + "/walk")
    assert page.locator("#step-2").count() == 1 and page.locator("#step-3").count() == 0
    for n in (3, 4, 5):
        nxt = page.locator(".walk-next a")
        assert nxt.count() == 1 and nxt.inner_text().startswith(f"Next: {n}")
        with page.expect_navigation():
            nxt.click()
        assert f"upto={n}" in page.url
        expect(page.locator(f"#step-{n}")).to_be_visible()
        expect(page.locator(f"#step-{n}")).to_be_in_viewport()  # the page scrolls to the new step
    assert page.locator(".walk-next a").count() == 0  # the last step has no "Next"


def test_feedback_button_sends_and_confirms(page):
    go(page, "/parts")
    page.locator(".fb-open").click()
    panel = page.locator("#fb-panel")
    assert panel.is_visible()
    expect(panel.locator(".fb-where")).not_to_have_text(re.compile(r"^(…)?$"))
    panel.locator("textarea[name=text]").fill("Browser test: the filter row could be shorter.")
    panel.get_by_role("button", name=re.compile("^Send")).click()
    expect(panel.locator(".fb-msg")).to_contain_text("Thank you")


# --- Systems & applications -------------------------------------------------------------------------------------------
def test_application_library_system_then_sub_system(page):
    go(page, "/blocks")
    system = page.locator(".sys-pick a.fchip").nth(1)
    name = system.inner_text().rsplit(" ", 1)[0].strip()
    system.click()
    page.wait_for_load_state("load")
    assert "system=" in page.url
    assert page.locator(".sys-pick a.fchip.on").inner_text().startswith(name)
    sub = page.locator(".sys-sub-title").first
    sub_name = sub.inner_text().strip()
    sub.click()
    page.wait_for_load_state("load")
    assert "sub=" in page.url
    heading = page.locator("h3.tpl-cat").first
    assert heading.is_visible() and heading.inner_text().startswith(sub_name)


def test_cross_sell_by_system(page):
    go(page, "/blocks/cross-sell")
    sel = page.locator("select[name=sel]")
    option = sel.locator("option").nth(2)  # the first sub-system of the first system
    with page.expect_navigation():
        sel.select_option(option.get_attribute("value"))
    assert "sel=" in page.url
    assert page.locator("main h2").first.is_visible()
    assert page.get_by_role("heading", name=re.compile("Our parts that go together")).is_visible()


def test_new_systems_in_the_libraries_and_cross_sell(page):
    """Owner, 25 Oct 2026: the diagrams under Other became new systems. The pages are opened again at 390 px by the
    fixture (no sideways scrolling)."""
    go(page, "/library")
    assert page.locator("#sys-aerospace").is_visible()
    go(page, "/blocks?system=aerospace")
    assert page.locator(".sys-pick a.fchip.on").inner_text().startswith("Aerospace")
    go(page, "/blocks?system=consumer&sub=connectivity")
    assert page.locator("#sub-connectivity").is_visible()
    go(page, "/blocks/cross-sell?sel=aerospace:satcom")
    assert page.get_by_role("heading", name=re.compile("Our parts that go together")).is_visible()
    go(page, "/blocks/cross-sell?sel=industrial")
    assert page.locator("main h2").first.is_visible()


def test_library_card_opens_its_block_diagram_in_one_click(page):
    """Owner, 25 Oct 2026: each card says whether a block diagram exists, and one click on a diagram name opens it
    (the editor is visited last: leaving it asks before unload)."""
    go(page, "/library")
    expect(page.locator(".sys-sub .dg-flag").first).to_be_visible()
    go(page, "/blocks")
    expect(page.locator("#diagrams")).to_be_in_viewport()  # the summary sits at the top
    card = page.locator(".sys-sub", has=page.locator(".dg-flag.has")).first
    expect(card.locator(".dg-flag")).to_contain_text("Block diagram ✓")
    with page.expect_navigation():
        card.locator(".sys-diags a").first.click()
    assert "/blocks/new?template=" in page.url


def test_efuse_key_applications_table_fits_desktop_and_phone(page, chromium, live_server):
    """Owner, 25 Oct 2026: the eFuse key-applications table: short cells without wrapping at 1280 px, cards at 390 px
    (no sideways scrolling of the page); "More" opens the details and the diagram count opens them too."""
    path = "/business/evidence/efuse_hotswap"
    page.set_viewport_size({"width": 1280, "height": 900})
    go(page, path)
    table = page.locator("table.ka-table")
    assert table.evaluate("t => t.scrollWidth <= t.parentElement.clientWidth + 1")
    heights = table.locator("tr.ka-main").evaluate_all("rows => rows.map(r => r.getBoundingClientRect().height)")
    assert max(heights) < 80, heights  # one or two lines per row, not a wall of wrapped text
    page.locator("a.ka-n").first.click()
    expect(page.locator(".ka-more details").first).to_have_attribute("open", "")
    ctx = chromium.new_context(viewport={"width": 390, "height": 844}, is_mobile=True)
    tab = ctx.new_page()
    tab.goto(live_server + path, wait_until="load")
    assert overflow(tab) <= 1
    box = tab.locator("table.ka-table").bounding_box()
    assert box["x"] >= 0 and box["x"] + box["width"] <= 391
    assert tab.locator("tr.ka-main th").first.evaluate("e => getComputedStyle(e).display") == "block"
    ctx.close()


def test_cross_sell_automotive_together_here_then_portfolio_tab(page):
    """Owner, 25 Oct 2026: pick Automotive on Cross-sell, see the Skyworks + Qorvo pairs there, open the portfolio tab
    with the system kept and its pairs marked, and come back to By system."""
    go(page, "/blocks/cross-sell")
    with page.expect_navigation():
        page.locator("select[name=sel]").select_option("automotive")
    block = page.locator("#together")
    assert block.get_by_role("heading", name=re.compile("Skyworks \\+ Qorvo together here")).is_visible()
    assert block.locator(".pf-syn-card").count() >= 5
    tab = page.locator("nav.page-tabs a", has_text="Skyworks + Qorvo portfolio")
    assert "system=automotive" in tab.get_attribute("href")
    with page.expect_navigation():
        tab.click()
    assert "/portfolio?system=automotive" in page.url
    assert page.locator(".pf-from").is_visible() and page.locator(".pf-syn-card.pf-here").count() >= 1
    with page.expect_navigation():
        page.locator("nav.page-tabs a", has_text="By system").click()
    assert "system=automotive" in page.url and page.locator("#together").is_visible()


def test_cross_sell_parts_table_and_automotive_filter(page):
    """Owner, 24-25 Oct 2026: the Skyworks and Qorvo parts table tops the selection; on Automotive the
    "Automotive-qualified only" filter is on (AEC-Q chips only, the hidden count shown), one click shows every part
    with its "not automotive" mark, and a pair link jumps to its card. Both views are checked at 390 px too."""
    go(page, "/blocks/cross-sell?system=automotive&sub=powertrain")
    table = page.locator("#parts-here")
    expect(table.get_by_role("heading", name=re.compile("Skyworks and Qorvo parts here"))).to_be_visible()
    toggle = page.get_by_role("button", name="Automotive-qualified only")
    expect(toggle).to_have_attribute("aria-pressed", "true")
    expect(table.locator(".aec-no")).to_have_count(0)
    with page.expect_navigation():
        toggle.click()
    assert "aec=0" in page.url
    expect(page.get_by_role("button", name="Automotive-qualified only")).to_have_attribute("aria-pressed", "false")
    expect(page.locator("#parts-here .aec-yes").first).to_be_visible()
    link = page.locator("#parts-here a.xs-pair").first
    if link.count():
        target = link.get_attribute("href")
        link.click()
        expect(page.locator(target)).to_be_in_viewport()


# --- Engineering knowledge --------------------------------------------------------------------------------------------
def test_rf_bands_type_and_application_chips_change_the_chart(page):
    go(page, "/rf-bands")
    chart = page.locator("#rbt-chart")
    chart.locator("g").first.wait_for(state="attached")
    before = chart.inner_html()
    chip = page.locator(".rbt-chip").first
    pressed = chip.get_attribute("aria-pressed")
    chip.click()
    assert chip.get_attribute("aria-pressed") != pressed
    after_type = chart.inner_html()
    assert after_type != before
    app = page.locator(".rbt-apps [data-app]").nth(1)
    app.click()
    assert app.get_attribute("aria-pressed") == "true"
    assert chart.inner_html() != after_type
    assert page.locator(".rbt-list").inner_text().strip()


def test_every_visual_moves_without_errors(page):
    go(page, "/visuals")
    links = page.locator("#vz-nav a[data-id]")
    ids = [links.nth(i).get_attribute("data-id") for i in range(links.count())]
    assert len(ids) > 10
    stuck, total = [], 0
    for vid in ids:
        page.locator(f"#vz-nav a[data-id='{vid}']").click()
        expect(page.locator("#vz-nav a.on")).to_have_attribute("data-id", vid)
        panel = page.locator("#vz-panel")
        before = panel.inner_html()
        moved = panel.evaluate("""p => {
          let n = 0;
          p.querySelectorAll('input[type=range]').forEach(r => {
            r.value = r.value === r.max ? r.min : r.max; r.dispatchEvent(new Event('input', {bubbles: true}));
            r.dispatchEvent(new Event('change', {bubbles: true})); n++;
          });
          p.querySelectorAll('select').forEach(s => {
            if (s.options.length < 2) return;
            s.selectedIndex = (s.selectedIndex + 1) % s.options.length;
            s.dispatchEvent(new Event('change', {bubbles: true})); s.dispatchEvent(new Event('input', {bubbles: true}));
            n++;
          });
          return n;
        }""")
        total += moved
        if moved and panel.inner_html() == before:
            stuck.append(vid)
    assert not stuck, f"controls moved but nothing changed in: {stuck}"
    assert total > len(ids)  # most visuals have several controls


def test_hex_visual_converts_flips_bits_and_follows_the_word_size(page):
    """Owner 24 Oct 2026: the hex / binary visual. Typing in one box fills the others, a clicked bit flips, the word
    size changes the signed reading, and a value too big for the word is refused."""
    go(page, "/visuals#hex")
    panel = page.locator("#vz-panel")
    expect(panel.locator("h2")).to_have_text("Hex, binary and registers")
    box = {k: panel.locator(f'[data-hx="{k}"]') for k in ("16", "2", "u", "s", "8")}
    box["16"].fill("0xFF")
    expect(box["u"]).to_have_value("255")
    expect(box["2"]).to_have_value("1111 1111")
    expect(box["s"]).to_have_value("-1")
    expect(panel.locator(".vz-kpis")).to_contain_text("0xFF")
    bit7 = panel.locator('.hx-bit[data-bit="7"]')
    expect(bit7).to_have_attribute("aria-pressed", "true")
    bit7.click()
    expect(panel.locator('.hx-bit[data-bit="7"]')).to_have_attribute("aria-pressed", "false")
    expect(box["16"]).to_have_value("0x7F")
    panel.locator('select[data-k="w"]').select_option("16")
    box["s"].fill("-1")
    expect(box["16"]).to_have_value("0xFFFF")
    expect(panel.locator(".hx-reg .hx-bit")).to_have_count(16)
    expect(panel.locator(".vz-more")).to_contain_text("0xFFFF")
    box["u"].fill("70000")
    expect(box["u"]).to_have_attribute("aria-invalid", "true")
    expect(panel.locator(".hx-msg")).to_contain_text("Does not fit in 16 bits")
    panel.locator(".hx-fields").fill("[15:8] HIGH, [3:0] LOW")
    expect(panel.locator(".hx-t").first).to_contain_text("HIGH")
    slider = panel.locator('input[data-k="bit"]')
    slider.evaluate("r => { r.value = 5; r.dispatchEvent(new Event('input', {bubbles: true})); }")
    expect(panel.locator(".vz-more")).to_contain_text("Worked example, bit 5")


def test_frequency_visuals_draw_the_bands_as_ranges(page):
    """Owner 24 Oct 2026: "FM is not only at 100 MHz". The wavelength chart shades AM, FM and DAB as ranges, the FM
    intermodulation chart shades the FM band by region, and the noise-figure visual shows the band strip."""
    go(page, "/visuals#wave")
    chart = page.locator("#vz-panel .vz-chart")
    for label in ("AM", "FM", "DAB", "GNSS", "UWB"):
        expect(chart.locator("text.vz-bl", has_text=label).first).to_be_attached()
    expect(page.locator("#vz-panel .vz-more")).to_contain_text("76–108 MHz")
    page.locator("#vz-nav a[data-id='fmim']").click()
    for label in ("Japan", "most countries", "OIRT"):
        expect(page.locator("#vz-panel .vz-chart text.vz-bl", has_text=label).first).to_be_attached()
    page.locator("#vz-nav a[data-id='nf']").click()
    expect(page.locator("#vz-panel .vz-strip .vz-sbar").first).to_be_attached()


# --- Data management --------------------------------------------------------------------------------------------------
def test_verify_next_answer_one_with_a_comment(page):
    go(page, "/manage/verify")
    assert page.title().startswith("Verify next") and "<" not in page.title()
    task = page.locator("article.vf-task").first
    title = task.locator("h3").inner_text()
    form = task.locator("form.vf-form").last
    form.locator("textarea[name=note]").fill("Checked in the browser test.")
    with page.expect_navigation():
        form.get_by_role("button", name="Correct").click()
    done = page.locator("details", has_text="Done").first
    done.locator("summary").click()
    assert page.locator("li", has_text=title).filter(has_text="Checked in the browser test.").count() >= 1


def test_overview_waiting_line_opens_the_filtered_requests(page):
    """Owner, 24 Oct 2026: a "Waiting for you" line goes straight to the place to act, here Requests > Part requests."""
    go(page, "/manage/requests")
    page.locator("details.rq-add summary").click()
    page.locator("textarea[name=mpns]").fill("ZQ4242")
    with page.expect_navigation():
        page.get_by_role("button", name="Add requests").click()
    go(page, "/manage/overview")
    line = page.locator("#ov-todo a", has_text="part request")
    with page.expect_navigation():
        line.click()
    assert page.url.endswith("/manage/requests?kind=part")
    expect(page.locator("a.fchip.on", has_text="Part requests")).to_be_visible()
    expect(page.locator("nav.page-tabs a.on", has_text="Requests")).to_be_visible()
    item = page.locator("li.rq-item", has_text="ZQ4242")
    expect(item).to_contain_text("Part request")
    with page.expect_navigation():
        item.locator("select[name=status]").select_option("in_progress")
    expect(page.locator("li.rq-item", has_text="ZQ4242").locator("select[name=status]")).to_have_value("in_progress")


# --- Shareable --------------------------------------------------------------------------------------------------------
def test_shareable_technology_overview_prints_to_four_a4_pages(page):
    """Shareable > Technology overview: the section row opens the landing page, its card the brief; the four sections
    show, the preview has no print button (owner, 26 Oct 2026: Create PDF makes the file instead), the page links to
    Create PDF, the print view still hides the site's frame, and the A4 print has 4 pages at most."""
    go(page, "/")
    page.locator("nav.main-nav a", has_text="Shareable").click()
    expect(page).to_have_url(re.compile(r"/shareable$"))
    expect(page.locator("main h1")).to_have_text("Shareable")
    # owner, 25 Oct 2026: one list; Preview opens the report in the lightbox, then check it as a page of its own
    page.locator(".sh-row", has_text="Technology overview").get_by_role("button", name="Preview").click()
    frame = page.locator("dialog.lb iframe.lb-frame")
    expect(frame).to_have_attribute("src", re.compile(r"/shareable/technology-overview\?embed=1$"))
    expect(page.frame_locator("dialog.lb iframe.lb-frame").locator("section#platforms")).to_be_attached()
    inner = page.frame_locator("dialog.lb iframe.lb-frame")
    expect(inner.locator("#tech-overview")).to_be_visible()
    for chrome in (".site-head", ".site-foot", ".to-actions"):  # only the report shows, like a document
        expect(inner.locator(chrome).first).to_be_hidden()
    expect(page.locator("dialog.lb").get_by_role("button", name=re.compile("Print"))).to_have_count(0)
    expect(page.locator("dialog.lb").get_by_role("link", name="Open in a new tab")).to_be_visible()
    page.get_by_role("button", name="Close preview").click()
    expect(page.locator("dialog.lb")).to_be_hidden()
    # tapping the report itself opens it too; ticking "Add to PDF" is separate (owner, 25 Oct 2026)
    page.locator(".sh-open", has_text="Skyworks + Qorvo synergies").click()
    expect(page.locator("dialog.lb")).to_be_visible()
    expect(frame).to_have_attribute("src", re.compile(r"/shareable/synergies\?embed=1$"))
    expect(page.locator('input[name="r"][value="synergies"]')).not_to_be_checked()
    page.get_by_role("button", name="Close preview").click()
    go(page, "/shareable/technology-overview")
    for sid, heading in (("at-a-glance", "technology overview"), ("platforms", "Technology platforms"),
                         ("markets", "End markets"), ("families", "Product families")):  # fmt: skip
        expect(page.locator(f"section#{sid} :is(h1, h2)").first).to_contain_text(heading)
    expect(page.locator(".to-matrix .to-dot").first).to_be_visible()
    expect(page.get_by_role("button", name="Create PDF")).to_have_attribute("href", "/shareable/pdf?r=overview&wm=none")
    expect(page.get_by_role("button", name="Save as PDF")).to_have_count(0)
    page.emulate_media(media="print")
    for frame in (".site-head", ".site-foot", ".fb", ".to-actions"):
        expect(page.locator(frame).first).to_be_hidden()
    expect(page.locator(".to-pfoot").first).to_contain_text("Public information")
    pdf = page.pdf(format="A4", print_background=True)
    pages = len(re.findall(rb"/Type\s*/Page\b(?!s)", pdf))
    assert 1 <= pages <= 4, f"{pages} A4 pages"


TO = "/shareable/technology-overview"


def pdf_pages(pdf: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page\b(?!s)", pdf))


def pdf_text(pdf: bytes, tmp_path, page_no: int | None = None) -> str | None:
    """The PDF's text (pdftotext from poppler), or None when the image has no pdftotext."""
    import shutil
    import subprocess

    if not shutil.which("pdftotext"):
        return None
    f = tmp_path / "page.pdf"
    f.write_bytes(pdf)
    pages = ["-f", str(page_no), "-l", str(page_no)] if page_no else []
    return subprocess.run(["pdftotext", *pages, str(f), "-"], capture_output=True, text=True, check=True).stdout


def a4(page) -> bytes:
    page.emulate_media(media="print")
    pdf = page.pdf(format="A4", print_background=True)
    page.emulate_media(media="screen")
    return pdf


def reset_overview(page):
    """The live server is shared by the browser tests: put the overview back to its defaults."""
    page.request.post(TO + "/edit/reset", form={"all": "1"})


def test_technology_overview_edit_number_paragraph_comment_resolve_undo(page, tmp_path):
    """Owner, 24 Oct 2026: edit a number, edit a paragraph, Save and see both (also in the PDF); comment on the
    paragraph, resolve the comment; Undo brings the page back."""
    page.on("dialog", lambda d: d.accept())
    try:
        go(page, TO)
        page.get_by_role("button", name="Edit", exact=True).click()
        expect(page).to_have_url(re.compile(r"/edit$"))
        expect(page.locator("#f-fact\\.parts\\.num .te-auto")).to_contain_text("from the catalogue")
        page.fill("input[name='fact.parts.num']", "5,000")
        page.locator("textarea[name='paras.glance']").first.fill("A new intro, written on the iPad.")
        page.get_by_role("button", name="Save").first.click()
        expect(page.locator(".te-status")).to_contain_text("Saved")
        expect(page.locator("#f-fact\\.parts\\.num .te-tag")).to_have_text("overridden")
        para = page.locator(".te-para").first
        para.locator("summary", has_text="Comment").click()
        para.locator("textarea[name^='ctext.']").fill("make this shorter")
        para.get_by_role("button", name="Send comment").click()
        expect(page.locator(".te-open").first).to_contain_text("make this shorter")
        page.locator(".te-open").first.get_by_role("button", name="Resolve").click()
        expect(page.locator(".te-status")).to_contain_text("Comment resolved")
        expect(page.locator(".te-open")).to_have_count(0)
        go(page, TO)
        expect(page.locator("#at-a-glance .to-lead")).to_have_text("A new intro, written on the iPad.")
        expect(page.locator(".to-fact b").first).to_have_text("5,000")
        text = pdf_text(a4(page), tmp_path)
        if text is not None:
            assert "A new intro, written on the iPad." in text and "5,000" in text
        go(page, TO + "/edit")
        page.get_by_role("button", name="Undo last change").click()
        go(page, TO)
        expect(page.locator(".to-fact b").first).not_to_have_text("5,000")
        expect(page.locator("#at-a-glance .to-lead")).not_to_contain_text("iPad")
    finally:
        reset_overview(page)


def test_technology_overview_is_four_pages_with_a_watermark(page, tmp_path):
    """Owner, 25 Oct 2026: the market focus page is gone, so the overview PDF is 4 A4 pages; an old ?focus= link opens
    the same 4 pages without error. The watermark prints on every page and changes no page count."""
    go(page, TO)
    expect(page.locator("select[name='focus']")).to_have_count(0)
    expect(page.locator("input[name='include']")).to_have_count(0)
    expect(page.locator(".to-page")).to_have_count(4)
    assert pdf_pages(a4(page)) == 4
    page.select_option("select[name='wm']", "draft")
    expect(page).to_have_url(re.compile(r"wm=draft"))
    expect(page.locator(".to-wm")).to_have_text("Draft")
    pdf = a4(page)
    assert pdf_pages(pdf) == 4
    for n in range(1, pdf_pages(pdf) + 1):
        text = pdf_text(pdf, tmp_path, n)
        if text is not None:
            assert "draft" in text.lower(), f"no watermark on page {n}"
    page.select_option("select[name='wm']", "none")
    expect(page.locator(".to-wm")).to_have_count(0)
    go(page, f"{TO}?focus=automotive")  # an old link: status 200, no error
    expect(page.locator("#market-focus")).to_have_count(0)
    expect(page.locator(".to-page")).to_have_count(4)
    assert pdf_pages(a4(page)) == 4
    go(page, f"{TO}?focus=bogus&wm=bogus")
    expect(page.locator(".to-wm")).to_have_count(0)


def test_technology_overview_logo_upload_and_remove(page):
    """Owner, 24 Oct 2026: add a logo (only with permission), see it on every page band, remove it."""
    from tests.test_tech_overview import png

    page.on("dialog", lambda d: d.accept())
    try:
        go(page, TO + "/edit")
        expect(page.locator(".te-logo")).to_contain_text("Use a logo only with permission from its owner.")
        page.set_input_files("input[name='logo']", files=[{"name": "logo.png", "mimeType": "image/png",
                                                           "buffer": png(40, 16)}])  # fmt: skip
        page.get_by_role("button", name="Add logo").click()
        expect(page.locator(".te-status")).to_contain_text("Logo added")
        go(page, TO)
        expect(page.locator(".to-band img.to-logo")).to_have_count(4)
        assert pdf_pages(a4(page)) <= 4
        go(page, TO + "/edit")
        page.get_by_role("button", name="Remove logo").click()
        expect(page.locator(".te-status")).to_contain_text("Logo removed")
        go(page, TO)
        expect(page.locator(".to-band img")).to_have_count(0)
    finally:
        reset_overview(page)


def _filler(n: int) -> str:
    words = "alpha bravo charlie delta echo foxtrot golf hotel india juliet kilo lima mike november oscar papa".split()
    out, i = "", 0
    while len(out) < n:
        out += ("" if not out else " ") + words[i % len(words)]
        i += 1
    return out[:n].rstrip()


def test_technology_overview_every_field_at_its_limit_still_fits(page):
    """The length limits keep the PDF at 4 A4 pages, even with every text and number at its
    limit, every dot set, two paragraphs per section and a medium logo."""
    from app.services import tech_overview_edit as te
    from tests.test_tech_overview import png

    form: dict = {}
    for fid, f in te.specs().items():
        if f.kind in ("line", "text", "data") and f.max and not fid.endswith(".url") and fid != "wm.custom":
            form[fid] = _filler(f.max)
        elif f.kind == "chips":
            form[fid] = "\n".join(_filler(te.CHIP_MAX) for _ in range(f.options[0] if f.options else te.CHIPS))
        elif f.kind in ("owner", "dot"):
            form[fid] = "Both"
    form["hdr.logo_size"] = "medium"
    multi = [(k, v) for k, v in form.items()]
    for s in te.sections():
        most, total = te.para_limits(s)
        multi += [(f"paras.{s}", _filler(total // most - 1)) for _ in range(most)] + [(f"pid.{s}", "")] * most
    try:
        r = page.request.post(TO + "/edit/logo", multipart={"logo": {"name": "l.png", "mimeType": "image/png",
                                                                      "buffer": png(40, 16)}})  # fmt: skip
        assert r.ok
        body = "&".join(f"{quote(k)}={quote(str(v))}" for k, v in multi)
        r = page.request.post(TO + "/edit", data=body, headers={"content-type": "application/x-www-form-urlencoded"})
        assert r.ok and "Not saved" not in r.text(), re.findall(r'role="alert">([^<]*)', r.text())[:5]
        go(page, TO)
        assert pdf_pages(a4(page)) <= 4
    finally:
        reset_overview(page)


def test_shareable_pack_of_all_three_and_each_report_alone(page):
    """Owner, 25 Oct 2026: the Synergies and Automotive reports are at most 3 A4 pages each; the combined PDF of all
    three reports at most 10 pages, numbered through. Owner, 26 Oct 2026: the PDF is made on the server; an old link
    to the combined print view downloads it."""
    for path in ("/shareable/synergies", "/shareable/automotive"):
        go(page, path)
        expect(page.locator(".to-page")).to_have_count(3)
        r = page.request.get(path.replace("/shareable/", "/shareable/pdf?r="))
        assert r.ok and r.headers["content-type"] == "application/pdf" and pdf_pages(r.body()) <= 3, path
    r = page.request.get("/shareable/pack?r=overview,synergies,automotive&wm=none")
    assert r.ok and r.headers["content-type"] == "application/pdf"
    assert pdf_pages(r.body()) <= 10


def test_shareable_create_pdf_with_two_reports(page, tmp_path):
    """The one Shareable page (owner, 26 Oct 2026): nothing ticked asks to pick a report; tick Synergies and
    Automotive, Create PDF downloads one A4 PDF (application/pdf) with Synergies then Automotive, made on the server;
    the page stays and the address keeps the choice."""
    go(page, "/shareable")
    box = page.locator("#build-pdf")
    expect(box.locator("input[name=r]:checked")).to_have_count(0)
    box.get_by_role("button", name="Create PDF").click()
    expect(page.locator("#sh-err")).to_have_text("Pick at least one report.")
    expect(page.locator("dialog.lb[open]")).to_have_count(0)
    box.locator("input[value='synergies']").check()
    box.locator("input[value='automotive']").check()
    expect(page.locator("#sh-err")).to_be_hidden()
    with page.expect_download() as dl:
        box.get_by_role("button", name="Create PDF").click()
    d = dl.value
    assert d.suggested_filename == "skyworks-qorvo-synergies-automotive.pdf"
    path = tmp_path / d.suggested_filename
    d.save_as(path)
    pdf = path.read_bytes()
    assert pdf.startswith(b"%PDF") and 2 <= pdf_pages(pdf) <= 6
    r = page.request.get("/shareable/pdf?r=synergies,automotive&wm=none")
    assert r.headers["content-type"] == "application/pdf"
    text = pdf_text(pdf, tmp_path)
    if text is not None:
        assert text.index("Skyworks + Qorvo synergies") < text.index("Extension of: Skyworks + Qorvo synergies")
        assert "Sources" in text and "BizDev" not in text
    expect(page).to_have_url(re.compile(r"/shareable\?r=synergies,automotive&wm=none$"))  # the page stays


def test_shareable_build_a_presentation_button(page):
    """Owner, 25 Oct 2026: Build a presentation sits next to Create PDF and opens the topic builder."""
    go(page, "/shareable")
    page.locator("#build-pdf").get_by_role("button", name="Build a presentation").click()
    expect(page).to_have_url(re.compile(r"/shareable/presentation$"), timeout=90_000)
    expect(page.locator("nav.sub-nav a.on")).to_have_text("Reports")


def test_shareable_create_powerpoint_with_two_reports(page, tmp_path):
    """Owner, 25 and 26 Oct 2026: Create PowerPoint sits with Create PDF; nothing ticked asks to pick a report; two
    ticked download one .pptx with both reports (iPad-size window)."""
    from pptx import Presentation

    go(page, "/shareable")
    box = page.locator("#build-pdf")
    box.get_by_role("button", name="Create PowerPoint").click()
    expect(page.locator("#sh-err")).to_have_text("Pick at least one report.")
    box.locator("input[value='synergies']").check()
    box.locator("input[value='automotive']").check()
    with page.expect_download() as dl:
        box.get_by_role("button", name="Create PowerPoint").click()
    d = dl.value
    assert d.suggested_filename == "skyworks-qorvo-synergies-automotive.pptx"
    path = tmp_path / d.suggested_filename
    d.save_as(path)
    heads = {sh.text_frame.text for s in Presentation(path).slides for sh in s.shapes
             if sh.has_text_frame and sh.text_frame.text.startswith("Skyworks + Qorvo · ")}  # fmt: skip
    assert heads == {"Skyworks + Qorvo · Synergies", "Skyworks + Qorvo · Automotive"}
    expect(page).to_have_url(re.compile(r"/shareable\?r=synergies,automotive&wm=none$"))  # the page stays


def test_shareable_synergies_owner_edit_and_undo(page):
    """The new reports edit like the overview: a number (the exact automatic value is shown), Save, Undo."""
    try:
        go(page, "/shareable/synergies")
        page.get_by_role("button", name="Edit", exact=True).click()
        expect(page.locator("#f-fact\\.parts\\.num .te-auto")).to_contain_text("(exact ")
        page.fill("input[name='fact.parts.num']", "4,600")
        page.get_by_role("button", name="Save").first.click()
        expect(page.locator(".te-status")).to_contain_text("Saved")
        go(page, "/shareable/synergies")
        expect(page.locator(".to-fact b").first).to_have_text("4,600")
        go(page, "/shareable/synergies/edit")
        page.get_by_role("button", name="Undo last change").click()
        go(page, "/shareable/synergies")
        expect(page.locator(".to-fact b").first).not_to_have_text("4,600")
    finally:
        page.request.post("/shareable/synergies/edit/reset", form={"all": "1"})


@pytest.mark.parametrize("slug", ["synergies", "automotive"])
def test_shareable_report_every_field_at_its_limit_still_fits(page, slug):
    """The length limits keep each management report at 3 A4 pages, with every text at its limit."""
    from app.services import share_reports as sr
    from app.services import tech_overview_edit as te

    rep = sr.REPORTS[slug]
    multi = []
    for fid, f in rep.specs().items():
        if f.kind in ("line", "text", "data") and f.max:
            multi.append((fid, _filler(f.max)))
        elif f.kind == "chips":
            multi += [(fid, "\n".join(_filler(te.CHIP_MAX) for _ in range(f.options[0] if f.options else te.CHIPS)))]
    for s in rep.sections():
        most, total = rep.para_limits(s)
        multi += [(f"paras.{s}", _filler(total // most - 1)) for _ in range(most)] + [(f"pid.{s}", "")] * most
    try:
        body = "&".join(f"{quote(k)}={quote(str(v))}" for k, v in multi)
        r = page.request.post(f"/shareable/{slug}/edit", data=body,
                              headers={"content-type": "application/x-www-form-urlencoded"})  # fmt: skip
        assert r.ok and "Not saved" not in r.text(), re.findall(r'role="alert">([^<]*)', r.text())[:5]
        go(page, f"/shareable/{slug}?wm=draft")
        assert pdf_pages(a4(page)) <= 3, slug
    finally:
        page.request.post(f"/shareable/{slug}/edit/reset", form={"all": "1"})


# --- Signing in -------------------------------------------------------------------------------------------------------
STUB = """
window.__calls = [];
addEventListener('click', () => { window.__inClick = true; }, true);
addEventListener('click', () => { window.__inClick = false; });
if (!window.PublicKeyCredential) window.PublicKeyCredential = function () {};
Object.defineProperty(navigator, 'credentials', {configurable: true, value: {
  get(o) { window.__calls.push({kind: 'get', inClick: !!window.__inClick,
                                challenge: !!(o && o.publicKey && o.publicKey.challenge)});
           return Promise.reject(new DOMException('cancelled in the test', 'NotAllowedError')); },
  create(o) { window.__calls.push({kind: 'create', inClick: !!window.__inClick});
              return Promise.reject(new DOMException('cancelled in the test', 'NotAllowedError')); },
}});
"""


def test_sign_in_passkey_button_calls_credentials_inside_the_tap(page):
    page.add_init_script(STUB)
    with page.expect_response(lambda r: r.url.endswith("/passkey/login/options")) as opts:
        go(page, "/login")
    assert opts.value.ok
    page.wait_for_timeout(50)  # the options are stored once the response is read
    page.locator("#passkey-login").click()
    calls = page.evaluate("window.__calls")
    assert calls == [{"kind": "get", "inClick": True, "challenge": True}]
    expect(page.locator("#auth-msg")).to_contain_text("cancelled")
    assert page.locator("#passkey-login").is_enabled()


# --- Research & ideas -------------------------------------------------------------------------------------------------
def test_expand_this_idea_asks_once_and_shows_the_request(page):
    """Owner, 25 Oct 2026: "Expand this idea" under a card opens a note field; sent, the button is gone on the list
    and on the report page, replaced by "Expansion requested" and the note (390 px checked after the test)."""
    go(page, "/business/reports")
    box = page.locator("details.rp-exp").first
    slug = box.get_attribute("id").removeprefix("exp-")
    expect(box.locator("summary")).to_contain_text("a deeper look at the applications before a business case")
    box.locator("summary").click()
    box.locator("textarea[name=note]").fill("Robot joints first")
    with page.expect_navigation():
        box.get_by_role("button", name="Ask for the expansion").click()
    done = page.locator(f"#exp-{slug}")
    expect(done).to_contain_text("Expansion requested")
    expect(done).to_contain_text("Robot joints first")
    assert page.locator(f"details#exp-{slug}").count() == 0
    go(page, f"/business/reports/{slug}")
    expect(page.locator(f"p#exp-{slug}")).to_contain_text("Robot joints first")
    assert page.locator("details.rp-exp").count() == 0
