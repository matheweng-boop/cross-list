"""Products › Cross a list (owner, 26 Oct 2026): a pasted list crossed in one go, two result columns (drop-in, best
functional fit), a traffic light per part, owner Right / Wrong answers, the Excel download and the accuracy check."""

import io
from datetime import date, timedelta
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select

from app.db import get_engine
from app.matching.engine import WORSE, Diff, Fit
from app.models import DataReview, ExternalSubmission
from app.services import crosslist as cl
from app.services import inbox
from app.services.seed import load_file
from tests.conftest import ROOT


def post(client, text: str):
    r = client.post("/cross-list", data={"parts": text})
    assert r.status_code == 200, r.text
    return r.text


def rows_for(session, text: str):
    entries, _ = cl.parse(text)
    return {r.text: r for r in cl.cross_list(session, entries)}


# --- parsing ---------------------------------------------------------------------------------------------------------
def test_parse_separators_excel_columns_and_duplicates():
    text = ("Part number\tQty\tPrice\n"  # a header row: no part number in it
            "SC141\t2000\t$0.85\tSynGamma\n"  # an Excel row: the first part-number cell only
            "SD4100, SA7741X; SC142\n"
            "\n"
            "sc-141\n"  # the same part again, written differently
            "1000 pcs, 10k, 0,85 €\n"
            "  ZZZ999  \n")  # fmt: skip
    entries, over = cl.parse(text)
    assert [e.text for e in entries] == ["SC141", "SD4100", "SA7741X", "SC142", "ZZZ999"] and over == 0
    assert entries[0].n == 2 and all(e.n == 1 for e in entries[1:])


def test_parse_keeps_the_first_500_lines():
    entries, over = cl.parse("\n".join(f"AB{i:04d}" for i in range(520)))
    assert len(entries) == cl.MAX_LINES and over == 20


# --- the page --------------------------------------------------------------------------------------------------------
def test_page_in_products_menu_with_privacy_note(client):
    page = client.get("/cross-list").text
    assert "Cross a list" in page and cl.PRIVACY_NOTE in page and 'name="parts"' in page
    assert "How the columns are chosen" in page and "Checked on priority-1 specs" in page
    assert 'href="/cross-list"' in client.get("/").text  # in the Products menu


@pytest.mark.smoke
def test_every_part_listed_in_order_with_lights_and_reasons(client):
    page = post(client, "SC141\nSD4100\nSC142\nZZZ999\nSC141\nSR8641\nSC14l")
    order = [page.index(f">{p}<") for p in ("SC141", "SD4100", "SC142", "ZZZ999")]
    assert order == sorted(order)
    assert "×2" in page  # SC141 twice, shown once
    assert "6 parts:" in page and "Coverage" in page and cl.PRECHECK_NOTE in page
    assert "Function known (" in page or "No Skyworks/Qorvo part for this function" in page  # SC142: no cross
    assert "Not in our data." in page  # ZZZ999
    assert "did you mean SC14" in page  # a typo
    assert "Already a SynRef part" in page  # our own part
    assert 'data-light="red"' in page and 'data-light="amber"' in page
    assert "Download Excel" in page


def test_rows_columns_and_comments(session, client):
    rows = rows_for(session, "SC141\nSD4100\nSC142\nZZZ999\nSC14I")
    dropin = rows["SC141"]
    assert dropin.drop_in.part.mpn == "SR8641" and dropin.drop_in.type == "Direct package equivalent"
    assert dropin.functional is None  # the best functional fit is the same part: shown once
    assert dropin.drop_in.comment.startswith("Same package and pinout")
    assert "demo data" in dropin.drop_in.comment and dropin.light == "amber"  # synthetic values: not from datasheets
    func = rows["SD4100"]
    assert func.drop_in is None and func.light == "amber"
    assert func.functional.type.startswith("Functional · other package")
    assert "Package SSOP-16 vs" in func.functional.comment
    red = rows["SC142"]
    assert red.light == "red" and red.drop_in is None and red.functional is None and red.reason
    assert rows["ZZZ999"].light == "red" and rows["ZZZ999"].reason == "Not in our data."
    assert rows["SC14I"].suggestion and "did you mean" in rows["SC14I"].reason


def test_ordering_code_variant_is_recognised(session):
    rows = rows_for(session, "SC141-TR\nSB8641")  # a reel suffix added; a suffix left out (SB8641N)
    assert rows["SC141-TR"].part.mpn == "SC141" and rows["SC141-TR"].how == "variant"
    assert rows["SB8641"].part.mpn == "SB8641N" and rows["SB8641"].how == "variant"


@pytest.mark.smoke
def test_known_drop_in_with_datasheets_is_green(session, real):
    rows = rows_for(session, "ISO6760QDWRQ1\nISO6431DBQR")
    g = rows["ISO6760QDWRQ1"]
    assert g.drop_in.part.mpn.startswith("SI86S660") and g.light == "green", g.drop_in.comment
    assert "AEC-Q both" in g.drop_in.comment and "(ours better)" in g.drop_in.comment
    a = rows["ISO6431DBQR"]
    assert a.light == "amber" and "(ours worse)" in (a.functional or a.drop_in).comment


def test_priority_2_does_not_demote_a_drop_in_priority_1_does(session, part):
    sc = part("SC141")
    pv = next(p for p in sc.params if p.key == "prop_delay_ns")
    pv.value, pv.num = 1.0, 1.0  # theirs much faster: ours worse on a priority 2 spec
    session.commit()
    r = rows_for(session, "SC141")["SC141"]
    assert r.drop_in.part.mpn == "SR8641" and r.drop_in.result.category == "fff"
    assert "Propagation delay" not in r.drop_in.comment  # priority 1 only
    cmti = next(p for p in sc.params if p.key == "cmti_kv_us")
    cmti.value, cmti.num = 500, 500.0  # theirs far higher: ours worse on priority 1
    session.commit()
    session.expire_all()
    r = rows_for(session, "SC141")["SC141"]
    assert r.drop_in is None or r.drop_in.result.category != "fff"
    col = r.drop_in or r.functional
    assert "CMTI" in col.comment and "(ours worse)" in col.comment and r.light == "amber"


# --- the column rules on hand-made results --------------------------------------------------------------------------
def _res(cid, mpn, fits=True, same_pkg=True, cat="fff", worse=0, close=0.0):
    devs = [Diff(f"k{i}", f"Spec {i}", 1, WORSE, "1", "2", 0.5) for i in range(worse)]
    cand = SimpleNamespace(id=cid, mpn=mpn, params=[])
    fit = Fit(same_pkg, "same" if same_pkg else "smaller", fits)
    return SimpleNamespace(candidate=cand, fit=fit, category=cat, deviations=devs, diffs=devs, exceptions=[],
                           closeness=close, lifecycle_risk=False)


def test_exact_equivalent_preferred_over_drop_in_with_deviations():
    ref = SimpleNamespace(params=[])
    exact = _res(1, "OURS-EXACT", cat="fff", close=50)
    near = _res(2, "OURS-NEAR", cat="package", worse=1, close=1)
    c1, c2 = cl.columns(ref, {"fff": [exact], "package": [near]})
    assert c1.candidate.mpn == "OURS-EXACT" and c2 is None  # the same part is never shown twice


def test_drop_in_with_fewest_deviations_and_functional_column():
    ref = SimpleNamespace(params=[])
    two = _res(1, "TWO-DEV", cat="package", worse=2)
    one = _res(2, "ONE-DEV", cat="package", worse=1)
    other_pkg = _res(3, "OTHER-PKG", fits=False, same_pkg=False, cat="functional", worse=0)
    c1, c2 = cl.columns(ref, {"package": [two, one], "functional": [other_pkg]})
    assert c1.candidate.mpn == "ONE-DEV" and c2.candidate.mpn == "OTHER-PKG"


def test_no_drop_in_only_a_functional_fit():
    ref = SimpleNamespace(params=[])
    f = _res(3, "OTHER-PKG", fits=False, same_pkg=False, cat="functional", worse=1)
    c1, c2 = cl.columns(ref, {"functional": [f]})
    assert c1 is None and c2.candidate.mpn == "OTHER-PKG"


# --- nothing stored, Excel -------------------------------------------------------------------------------------------
def test_nothing_is_written_to_the_database(client):
    writes = []

    def watch(conn, cursor, statement, *a):
        if statement.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")):
            writes.append(statement)

    event.listen(get_engine(), "before_cursor_execute", watch)
    try:
        post(client, "SC141\nSD4100\nZZZ999")
        client.post("/cross-list.xlsx", data={"parts": "SC141\nSD4100"})
    finally:
        event.remove(get_engine(), "before_cursor_execute", watch)
    assert writes == []


def test_excel_download_mirrors_the_table(client):
    r = client.post("/cross-list.xlsx", data={"parts": "SC141\nSD4100\nSC142\nZZZ999\nSC141"})
    assert r.status_code == 200
    assert f"cross-list-{date.today():%Y-%m-%d}.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Cross list"]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0] == cl.EXPORT_HEADER and ws["A1"].font.bold and ws.freeze_panes == "A2" and ws.auto_filter.ref
    body = values[1:]
    assert [v[1] for v in body] == ["SC141 ×2", "SD4100", "SC142", "ZZZ999"]
    assert [v[4] for v in body] == ["Amber", "Amber", "Red", "Red"]
    assert [ws.cell(i, 5).fill.fgColor.rgb[-6:] for i in range(2, 6)] == ["FFEB9C", "FFEB9C", "FFC7CE", "FFC7CE"]
    assert body[0][5] == "SR8641" and body[0][7] == "Direct package equivalent" and body[0][9] == "= Drop-in"
    assert body[0][8].startswith("Same package and pinout") and ws.cell(2, 9).alignment.wrap_text
    assert body[1][5] == "No drop-in" and body[1][11].startswith("Functional")
    assert body[2][9] == "No functional fit" and body[2][12]  # the reason
    assert body[3][2] == "not recognised" and body[3][12] == "Not in our data."
    assert body[0][13] is None  # no other Skyworks isolator family option for synthetic parts
    assert "Isolation rating" in body[0][14] and body[0][16] is None  # checked on; the owner's comment column
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)}
    assert summary["Parts"] == 4 and summary["Red"] == 2 and summary["Coverage"] == "50 %"
    assert cl.PRECHECK_NOTE in [r[1] for r in wb["Summary"].iter_rows(values_only=True)]
    shown = client.post("/cross-list.xlsx", data={"parts": "SC141\nSD4100\nSC142\nZZZ999", "lights": "red"})
    rows = list(load_workbook(io.BytesIO(shown.content))["Cross list"].iter_rows(values_only=True))
    assert [r[1] for r in rows[1:]] == ["SC142", "ZZZ999"]


def test_copy_table_text_has_a_header_and_one_line_per_part(session):
    entries, _ = cl.parse("SC141\nSD4100\nZZZ999\nSC141")
    lines = cl.tsv(cl.cross_list(session, entries)).splitlines()
    assert lines[0].split("\t") == cl.EXPORT_HEADER and len(lines) == 4
    assert all(len(x.split("\t")) == len(cl.EXPORT_HEADER) for x in lines)


def test_page_has_the_copy_table_source(client):
    page = post(client, "SC141\nZZZ999")
    assert 'id="cx-copy"' in page and 'id="cx-export"' in page and "Export shown rows" in page
    assert page.count("<tr data-light=") == 2


def test_over_500_lines_message(client):
    page = post(client, "\n".join(f"AB{i:04d}" for i in range(505)))
    assert "Only the first 500 lines were crossed; 5 more lines left out" in page


# --- Right / Wrong (owner only), verified everywhere, inbox ---------------------------------------------------------
def test_owner_marks_right_and_wrong(owner, session):
    page = owner.post("/cross-list", data={"parts": "SC141\nSD4100"}).text
    assert 'name="verdict" value="right"' in page
    r = owner.post("/cross-list/feedback", data={"their": "SC141", "ours": "SR8641", "verdict": "right"})
    assert r.status_code == 200 and "Verified" in r.text
    need = owner.post("/cross-list/feedback", data={"their": "SD4100", "ours": "SR8640N", "verdict": "wrong"})
    assert "Say what is wrong" in need.text  # a Wrong needs a reason; nothing stored
    owner.post("/cross-list/feedback", data={"their": "SD4100", "ours": "SR8640N", "verdict": "wrong",
                                             "comment": "isolation too low for the customer"})  # fmt: skip
    got = session.scalars(select(DataReview).where(DataReview.dataset == cl.DATASET)).all()
    assert [(x.key, x.verdict) for x in got] == [("SC141|SR8641", "right"), ("SD4100|SR8640N", "wrong")]
    page = owner.post("/cross-list", data={"parts": "SC141\nSD4100"}).text
    assert "✓ Verified" in page and "Marked wrong" in page
    sc = owner.get("/find?q=SC141", follow_redirects=True).text  # the single search shows it too
    assert "✓ Verified" in sc
    text = inbox.collect(session, date.today() - timedelta(days=1))
    wrong = text.split("## Cross list: wrong crosses")[1].split("##")[0]
    assert "SD4100 -> SR8640N · isolation too low for the customer" in wrong
    assert "SC141 -> SR8641" in text.split("## Cross list: confirmed crosses")[1].split("##")[0]
    assert "crosscheck" not in text.split("## Review answers")[1].split("##")[0]


def test_guest_sees_results_without_buttons_and_cannot_answer(make_guest, session):
    guest, _ = make_guest()
    page = guest.post("/cross-list", data={"parts": "SC141\nZZZ999"})
    assert page.status_code == 200 and "Not in our data." in page.text and 'value="right"' not in page.text
    assert guest.post("/cross-list.xlsx", data={"parts": "SC141"}).status_code == 200
    assert session.scalar(select(ExternalSubmission.id)) is None  # compute only, never queued for review
    r = guest.post("/cross-list/feedback", data={"their": "SC141", "ours": "SR8641", "verdict": "right"})
    assert r.status_code == 403
    assert session.scalar(select(DataReview.id).where(DataReview.dataset == cl.DATASET)) is None


# --- daily accuracy check --------------------------------------------------------------------------------------------
def test_accuracy_script(session, tmp_path, monkeypatch):
    from scripts import check_cross_accuracy as script

    monkeypatch.setattr(cl, "known_pairs", lambda: [
        {"their": "SC141", "ours": "SR8641", "expect": "right"},  # shown: a hit
        {"their": "SD4100", "ours": "SR8621", "expect": "right"},  # not shown: a miss
        {"their": "SC141", "ours": "SR8640N", "expect": "wrong"},  # not shown: a hit
        {"their": "NOPE123", "ours": "SR8641", "expect": "right"},  # not in the data: left out
    ])  # fmt: skip
    out = tmp_path / "acc.yaml"
    res = script.main(["--write"], session=session, out=out)
    assert (res["n"], res["hits"], res["rate"], res["not_in_data"]) == (3, 2, 67, ["NOPE123"])
    from app import yamlio

    saved = yamlio.load(out.read_text())
    assert saved["rate"] == 67 and saved["n"] == 3 and len(saved["misses"]) == 1


def test_known_pairs_include_engineer_notes():
    pairs = cl.known_pairs()
    assert any(p["source"].startswith("engineer note") for p in pairs)


def test_page_shows_the_accuracy_line_once_enough_pairs_are_known(client, monkeypatch):
    monkeypatch.setattr(cl, "accuracy", lambda: {"rate": 88, "n": 25, "checked": "2026-10-26"})
    assert "Drop-ins correct" not in client.get("/cross-list").text  # hidden below 30 known pairs (owner)
    monkeypatch.setattr(cl, "accuracy", lambda: {"rate": 93, "n": 30, "checked": "2026-10-26"})
    assert "Drop-ins correct: <b>93 %</b> of 30 known pairs (checked 2026-10-26)" in client.get("/cross-list").text


def test_p1_specs_behind_a_tap_and_no_repeated_reason(client):
    html = client.post("/cross-list", data={"parts": "XYZ12345"}).text
    cell = html[html.index('class="cx-reason"'):]
    cell = cell[:cell.index("</td>")]
    assert cell.count("Not in our data") == 1  # said once in the row (the copy-table text keeps its own)
    ok = client.post("/cross-list", data={"parts": "ISO7741DW"}).text
    if "P1 specs checked" in ok:
        assert '<details class="cx-p1"><summary>P1 specs checked' in ok


@pytest.mark.parametrize("tok,ok", [("SI8641BD-IS", True), ("2000", False), ("$0.85", False), ("10k", False),
                                    ("Texas", False), ("ISO7741DWR", True), ("AB12", True)])  # fmt: skip
def test_looks_like_pn(tok, ok):
    assert cl.looks_like_pn(tok) is ok


def _fake(mpn: str, pid: int, *, fff: bool, fits: bool, same_package: bool = True):
    from types import SimpleNamespace as NS

    return NS(candidate=NS(id=pid, mpn=mpn, params=[]), category="fff" if fff else "functional", deviations=[],
              exceptions=[], closeness=0, diffs=[], lifecycle_risk=False,
              fit=NS(fits=fits, same_package=same_package, same_pinout=same_package))  # fmt: skip


def test_columns_keep_the_best_functional_fit_for_an_si86s_drop_in():
    """The Si86S -> Si86 rule (owner, 26 Oct 2026) moved out of column 2 into the other-family option (27 Oct)."""
    from types import SimpleNamespace as NS

    s = _fake("SI86S640EE-IS2", 1, fff=True, fits=True)
    other = _fake("SI8285DD-IS", 2, fff=False, fits=False, same_package=False)
    si86 = _fake("SI8641ET-IS", 3, fff=False, fits=False, same_package=False)
    si86.deviations = [NS(score=1.0)]  # ranks below `other` as a plain functional fit
    assert cl.columns(NS(params=[]), {"x": [s, other, si86]}) == (s, None)  # "= Drop-in", no Si86 swap
    assert cl.is_si86s("Si86S641HE-IS2") and cl.is_si86("SI8641ET-IS") and not cl.is_si86("SI86SL41EC-IS1")
    assert [cl.isolator_family(m) for m in ("SI86S641EC-IS1", "SI8641BD-IS", "SI8285DD-IS")] == ["Si86S", "Si86", None]


def test_other_family_both_directions():
    """Owner, 27 Oct 2026: "always offer both families if both families are available" - Si86S shown: the best
    suitable Si86 part; Si86 shown: the best suitable Si86S part; by functional_key, parts shown left out."""
    from types import SimpleNamespace as NS

    s = _fake("SI86S640EE-IS2", 1, fff=True, fits=True)
    other = _fake("SI8285DD-IS", 2, fff=False, fits=False, same_package=False)
    si86 = _fake("SI8641ET-IS", 3, fff=False, fits=False, same_package=False)
    si86_worse = _fake("SI8641BD-IS", 4, fff=False, fits=False, same_package=False)
    si86_worse.deviations = [NS(score=1.0)]
    pool = [s, other, si86_worse, si86]
    assert cl.other_family([s, other], pool) == ("Si86", si86)  # an Si86S drop-in still gets an Si86 option
    assert cl.other_family([None, s], pool) == ("Si86", si86)  # no drop-in: the functional fit alone
    assert cl.other_family([si86, None], pool) == ("Si86S", s)  # an Si86 drop-in gets an Si86S option
    assert cl.other_family([s, si86], pool) is None  # both families shown already
    assert cl.other_family([other, None], pool) is None  # not a Skyworks isolator family
    assert cl.other_family([s, other], [s, other]) is None  # no suitable part in the other family
    assert cl.other_family([si86, None], [si86, si86_worse, other]) is None


PRELOAD = {"isolators": (("novosense_digital_isolators", "skyworks_si86", "skyworks_isolators_more"), None)}


@pytest.fixture
def isolators(session):
    """The Novosense and Skyworks digital isolator files in the test database (PRELOAD, tests/conftest.py)."""


def test_nsi8321_offers_both_families_on_cross_a_list(session, isolators):
    """Owner, 27 Oct 2026: NSI8321N1-Q1SPR gave one family only. No drop-in; the functional fit is Si86S, the Si86
    family is offered under it (export and Copy table too)."""
    row = rows_for(session, "NSI8321N1-Q1SPR")["NSI8321N1-Q1SPR"]
    assert row.drop_in is None and cl.isolator_family(row.functional.part.mpn) == "Si86S"
    assert row.other is not None and cl.isolator_family(row.other.part.mpn) == "Si86"
    assert row.other.type == "Also in the Si86 family" and "CMTI" in row.other.comment
    cell = cl.export_cells(row)[cl.EXPORT_HEADER.index("Other family option")]
    assert cell.startswith(f"Also in the Si86 family: {row.other.part.mpn} – ")
    assert row.other.part.mpn in cl.tsv([row])


def test_other_family_in_the_excel_and_on_the_page(client, isolators):
    r = client.post("/cross-list.xlsx", data={"parts": "NSI8321N1-Q1SPR"})
    ws = load_workbook(io.BytesIO(r.content))["Cross list"]
    head = [c.value for c in ws[1]]
    assert "Other family option" in head
    assert ws.cell(2, head.index("Other family option") + 1).value.startswith("Also in the Si86 family: SI86")
    assert "Also in the Si86 family" in post(client, "NSI8321N1-Q1SPR")


def test_part_page_shows_the_other_family(client, session, isolators):
    """The part page ranks by the engine's order (Si86 first for NSI8321N1-Q1SPR): its Other family line offers the
    best Si86S part, so both families show there as well."""
    from app.models import Part

    pid = session.scalar(select(Part.id).where(Part.mpn == "NSI8321N1-Q1SPR"))
    html = client.get(f"/parts/{pid}").text
    line = html[html.index('class="alt-line other-family"'):]
    line = line[:line.index("</p>")]
    assert "Other family:</b> also in the Si86S family" in line and "SI86S6" in line and "CMTI" in line
    ours = session.scalar(select(Part.id).where(Part.mpn == "SI8641BD-IS"))
    assert ours and 'class="alt-line other-family"' not in client.get(f"/parts/{ours}").text  # nothing crossed: no line


# --- TI automotive reel codes and the isolated DC-DC family (owner, 26 Oct 2026) --------------------------------------
def test_reel_letter_before_q1_is_a_variant():
    from app.partnumbers import reel_q1_variants

    assert reel_q1_variants("ISOW7841AQDWERQ1") == ["ISOW7841AQDWEQ1"]
    assert reel_q1_variants("ISOW7841AQDWEQ1") == ["ISOW7841AQDWERQ1"]
    assert reel_q1_variants("SI8641BD") == [] and reel_q1_variants("RQ1") == []


def test_ti_reel_code_before_q1_is_recognised(session, real):
    from app.services.queries import search_parts

    r = rows_for(session, "ISOW7841AQDWERQ1")["ISOW7841AQDWERQ1"]
    assert r.part.mpn == "ISOW7841AQDWEQ1" and r.how == "variant"
    assert "Function known (Digital isolator)" in r.reason  # no Skyworks isolator with an integrated-power module
    assert [p.mpn for p in search_parts(session, "ISOW7841AQDWERQ1")] == ["ISOW7841AQDWEQ1"]  # the single search too


def test_isolated_dcdc_module_is_known_without_a_cross(session):
    load_file(session, ROOT / "data" / "real" / "ti_isolated_dcdc.yaml")
    session.commit()
    rows = rows_for(session, "UCC12051QDVERQ1\nUCC12051QDVEQ1")
    for text, how in (("UCC12051QDVERQ1", "exact"), ("UCC12051QDVEQ1", "variant")):
        r = rows[text]
        assert r.part.mpn == "UCC12051QDVERQ1" and r.how == how and r.part.family == "isolated_dcdc"
        assert r.light == "red" and r.drop_in is None and r.functional is None
        assert r.reason == ("Function known (Isolated DC-DC, integrated-transformer module), no Skyworks/Qorvo part "
                            "in our data.")  # fmt: skip
        assert r.reason_link == "/business/evidence/isolated_dcdc"
    p = {pv.key: pv.value for pv in rows["UCC12051QDVERQ1"].part.params}
    assert p["isolation_rating_kvrms"] == 5.0 and p["qualification"] == "automotive" and p["output_power_mw"] == 500.0
