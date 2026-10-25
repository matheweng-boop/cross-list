"""Presentation builder > Download as PowerPoint (owner, 25 Oct 2026): the .pptx holds the preview's slides in the
same order with the same title, the parts slide is a native table, every text passes the site word check, the button
is on the builder and the preview, and the route keeps the preview's access rules."""

import io

from pptx import Presentation

from app.services import presentation as pr
from app.services import presentation_pptx as pp
from app.services import site_guard
from tests.test_presentation import UWB

URL = "/shareable/presentation/pptx?s=" + UWB


def _open(data: bytes):
    return Presentation(io.BytesIO(data))


def test_download_is_a_pptx_of_the_preview_slides(client):
    r = client.get(URL)
    assert r.status_code == 200
    assert r.headers["content-type"] == pp.MEDIA_TYPE
    assert r.headers["content-disposition"] == 'attachment; filename="skyworks-qorvo.pptx"'
    assert r.headers.get("x-content-type-options") == "nosniff"  # the normal security headers
    slides = pr.build(pr.parse(UWB))
    prs = _open(r.content)
    assert prs.slide_width * 9 == prs.slide_height * 16  # 16:9
    assert len(prs.slides) == len(slides)  # title + one slide per preview slide of each section, same order
    texts = pp.texts(r.content)
    for want, got in zip(slides, texts, strict=True):
        assert "Skyworks + Qorvo · Skyworks + Qorvo" in got  # the header, as in the preview
        assert any("Public information" in t for t in got)
        if want["kind"] != "title":
            assert want["title"] in got
    assert "Skyworks + Qorvo" in texts[0]  # the title slide heading
    kinds = [s["kind"] for s in slides]
    table_slide = prs.slides[kinds.index("parts")]
    assert any(sh.has_table for sh in table_slide.shapes)
    assert "Skyworks + Qorvo · Skyworks + Qorvo" == prs.core_properties.title


def test_every_text_in_the_file_passes_the_site_word_check(client):
    data = client.get(URL + "&wm=draft").content
    texts = [t for slide in pp.texts(data) for t in slide]
    assert any("DRAFT" in t for t in texts) and any("Public information · Draft" in t for t in texts)
    assert [h for t in texts for h in site_guard.hits(t)] == []
    assert all(pr.safe(t) for t in texts if t not in ("DRAFT",))


def test_every_section_builds_and_stays_site_clean():
    data = pp.build(pr.build(pr.all_section_ids()), pr.DEFAULT_TITLE)
    assert [h for slide in pp.texts(data) for t in slide for h in site_guard.hits(t)] == []


def test_saved_deck_title_names_the_file_and_guests_get_the_same_access(make_guest, locked):
    from tests.conftest import sign_in_owner

    guest, _ = make_guest()
    assert guest.get(URL).status_code == 200
    assert sign_in_owner(locked).status_code == 303
    form = {"name": "UWB for a car maker", "s": UWB, "title": "UWB in the car"}
    assert locked.post("/shareable/presentation/save", data=form, follow_redirects=False).status_code == 303
    r = locked.get("/shareable/presentation/pptx?deck=uwb-for-a-car-maker")
    assert r.headers["content-disposition"] == 'attachment; filename="uwb-in-the-car.pptx"'
    assert "Skyworks + Qorvo · UWB in the car" in pp.texts(r.content)[1]
    assert "/shareable/presentation/pptx?deck=uwb-for-a-car-maker" in locked.get("/shareable/presentation").text


def test_route_is_get_only_and_empty_goes_back(client):
    assert client.post(URL).status_code == 405
    r = client.get("/shareable/presentation/pptx", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/shareable/presentation")


def test_button_on_the_builder_and_the_preview(client):
    b = client.get("/shareable/presentation?s=" + UWB).text
    assert "Download as PowerPoint" in b and 'id="pb-pptx"' in b and "/shareable/presentation/pptx?s=" in b
    p = client.get("/shareable/presentation/preview?s=" + UWB).text
    assert "Download as PowerPoint" in p and "/shareable/presentation/pptx?s=" in p


def test_filename_is_safe():
    assert pp.filename('a/../"b" ü') == "a-b.pptx" and pp.filename("") == "presentation.pptx"
