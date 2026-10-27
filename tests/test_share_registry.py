"""Shareable reports come from one registry (share_reports.REGISTRY, owner 25 Oct 2026: "if I create a drones report,
then I would want a third page to show up ... and then be able to select if I want to merge these"), and Download as
PowerPoint makes one .pptx of the ticked reports. The browser click is in tests/test_browser.py."""

import dataclasses
import io
import re
from html import unescape

import pytest
from pptx import Presentation

from app.services import presentation_pptx as pp
from app.services import share_doc as sd
from app.services import share_pdf as spdf
from app.services import share_reports as sr
from app.services import site_guard

PPTX = "/shareable/pptx"


@pytest.fixture
def drones(monkeypatch):
    """A fake extra report, built like Automotive (same data and body), added to the registry for this test only."""
    rep = dataclasses.replace(sr.AUTOMOTIVE, key="drones", page="share_drones", label="Drones: test report",
                              url="/shareable/drones")  # fmt: skip

    def slides(data, vals, today=None, ip_note=""):
        return [{"kind": "title", "head": "Drones", "title": "Drones", "body": {"heading": "Drones", "date": today}},
                *sr.ss.table("Drones", "Drones table", ["A", "B"], [["1", "2"]])]  # fmt: skip

    def doc(data, vals, today=None):
        return {**sd.automotive(data, vals, today), "key": "drones", "title": "Drones: test report", "head": "Drones"}

    entry = sr.ShareReport("drones", "Drones", "A test report.", rep.url, 2, sr.build_automotive, slides,
                           "shareable/_automotive.html", rep, doc=doc)  # fmt: skip
    monkeypatch.setattr(sr, "REGISTRY", [*sr.REGISTRY, entry])
    return entry


def rows(html: str) -> list[str]:
    return re.findall(r'<li class="sh-row">.*?</li>', html, re.S)


def test_an_extra_registry_entry_shows_as_a_row_in_the_pack_and_in_powerpoint(client, drones):
    html = client.get("/shareable").text
    assert len(rows(html)) == len(sr.REGISTRY)  # the reports and the system briefs, then drones
    last = unescape(rows(html)[-1])
    assert 'name="r" value="drones"' in last and "<b>Drones</b>" in last and "2 pages" in last
    assert 'class="sh-open" href="/shareable/drones" data-lightbox="Drones"' in last and "Include" in last
    assert client.get("/shareable/drones").status_code == 200  # its own page, via the registry
    pdf = client.get("/shareable/pdf?r=drones,synergies")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == spdf.MEDIA_TYPE
    pages = spdf.texts(pdf.content)
    assert "Skyworks + Qorvo synergies" in pages[0] and "Drones: test report" in "".join(pages)
    assert f"page {len(pages)} / {len(pages)}" in pages[-1] and len(pages) <= 5
    r = client.get(f"{PPTX}?r=drones,synergies")
    assert r.status_code == 200
    texts = pp.texts(r.content)
    heads = [next(x for x in t if x.startswith("Skyworks + Qorvo · ")) for t in texts]
    assert heads[0] == "Skyworks + Qorvo · Synergies" and heads[-1] == "Skyworks + Qorvo · Drones"  # list order


def test_pptx_has_the_ticked_reports_in_list_order_with_tables_and_watermark(client):
    r = client.get(f"{PPTX}?r=overview,automotive,synergies&wm=draft")
    assert r.status_code == 200 and r.headers["content-type"] == pp.MEDIA_TYPE
    assert 'attachment; filename="skyworks-qorvo-synergies-automotive-overview.pptx"' == r.headers[
        "content-disposition"]  # fmt: skip
    prs = Presentation(io.BytesIO(r.content))
    assert prs.slide_width * 9 == prs.slide_height * 16  # 16:9
    texts = pp.texts(r.content)
    heads = [next(x for x in t if x.startswith("Skyworks + Qorvo · ")) for t in texts]
    order = list(dict.fromkeys(heads))
    assert order == ["Skyworks + Qorvo · Synergies", "Skyworks + Qorvo · Automotive",
                     "Skyworks + Qorvo · Technology overview"]  # fmt: skip
    for h in order:
        assert heads.count(h) >= 4  # cover, at a glance, tables, sources
    tables = [sh for s in prs.slides for sh in s.shapes if sh.has_table]
    assert len(tables) >= 8
    every = [x for t in texts for x in t]
    assert "DRAFT" in every and any("Public information" in x for x in every)
    assert "Combined at a glance" in every and "Where the two combine" in every and "New technology" in every
    assert not [h for x in every for h in site_guard.hits(x)]  # the public-safe word check
    # the text stays editable: titles are text boxes, not pictures
    assert not [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 13]  # 13: picture (no logo set)


def test_pptx_of_one_report_and_the_owners_edits(client):
    client.post("/shareable/automotive/edit", data={"title": "Cars: our edited title"})
    try:
        r = client.get(f"{PPTX}?r=automotive")
        texts = pp.texts(r.content)
        assert "Cars: our edited title" in texts[0]
        assert all(any(x == "Skyworks + Qorvo · Automotive" for x in t) for t in texts)
    finally:
        client.post("/shareable/automotive/edit/reset", data={"all": "1"})


def test_pptx_with_nothing_ticked_goes_back_and_asks_to_pick(client):
    for q in ("", "?r=", "?r=nothing"):
        r = client.get(PPTX + q, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/shareable?pick=1#build-pdf"
    html = client.get("/shareable?pick=1").text
    assert '<p class="sh-err" id="sh-err" role="alert">Pick at least one report.</p>' in html
    assert 'id="sh-err" role="alert" hidden>' in client.get("/shareable").text


def test_pptx_refuses_when_the_word_check_finds_something(client, monkeypatch):
    monkeypatch.setattr(site_guard, "hits", lambda text: ["X"] if "Synergies" in text else [])
    assert client.get(f"{PPTX}?r=synergies").status_code == 409


def test_create_pdf_and_powerpoint_sit_with_the_main_buttons(client):
    html = client.get("/shareable").text
    btns = html.split('class="sh-btns"', 1)[1].split("</span>", 1)[0]
    assert btns.index(">Create PDF<") < btns.index(">Create PowerPoint<") < btns.index(">Build a presentation<")
    assert 'id="sh-pptx" data-href="/shareable/pptx"' in btns


def test_every_report_editor_is_owner_only_also_a_new_one():
    from app.auth.access import owner_only

    for path in ("/shareable/drones/edit", "/shareable/drones/edit/undo", "/shareable/synergies/edit"):
        assert owner_only(path, "GET") and owner_only(path, "POST"), path
    assert not owner_only("/shareable/drones", "GET") and not owner_only("/shareable/pptx", "GET")


def test_registry_keeps_the_list_and_pack_orders():
    briefs = ["uwb", "v2x", "gnss", "radio"]  # the system briefs (owner, 27 Oct 2026), tests/test_share_systems.py
    assert [r.key for r in sr.REGISTRY] == ["synergies", "automotive", "overview", *briefs]
    assert sr.pack_order() == ("overview", "synergies", "automotive", *briefs) == sr.PACK_ORDER
    assert sr.list_order(["overview", "x", "synergies", "synergies"]) == ["synergies", "overview"]
    assert set(sr.reports()) == {"synergies", "automotive", *briefs} and sr.pages() == sr.PAGES
