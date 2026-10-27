import csv
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from tests.conftest import ids as find_id


def ids(client, *mpns):
    out = []
    for m in mpns:
        out.append(next(p["id"] for p in client.get(f"/api/parts?q={m}").json() if p["mpn"] == m))
    return out


def test_pages_render(client):
    for url in ["/parts", "/parts?role=&q=SA", "/compare", "/rules"]:
        r = client.get(url)
        assert r.status_code == 200, url
        assert "synthetic" in r.text.lower()


def test_part_page_shows_cross_types_and_top5(client):
    (ref,) = ids(client, "SR8641")
    html = client.get(f"/parts/{ref}").text
    for label in [
        "Drop-in",
        "Functional",
        "Package eq.",
        "Pin-to-pin",
        "Electrical (priority 1)",
        "Top 5 differences",
        "Download table (Excel)",
        "compare-tray",
        "Qualification",
    ]:
        assert label in html, label


def test_hide_eol_filter(client):
    (ref,) = ids(client, "SR8641")
    assert "SA7741X" in client.get(f"/parts/{ref}?all=true").text
    assert "SA7741X" not in client.get(f"/parts/{ref}?hide_eol=true&all=true").text


def test_missing_part_404(client):
    assert client.get("/parts/99999").status_code == 404


def test_cross_excel_highlights_differences(client):
    ref, cand = ids(client, "SR8641", "SB8441")
    r = client.get(f"/cross/{ref}/{cand}.xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = {row[1].value: row for row in ws.iter_rows() if len(row) >= 5 and row[1].value}
    default_out = rows["Output type (default state)"]
    assert default_out[4].value == "different"
    assert default_out[3].fill.fgColor.rgb.endswith("FFEB9C")


def test_cross_csv(client):
    ref, cand = ids(client, "SR8641", "SA7741")
    rows = list(csv.reader(io.StringIO(client.get(f"/cross/{ref}/{cand}.csv").text)))
    assert rows[0][0] == "Priority"
    iso = next(r for r in rows if r[1] == "Isolation rating")
    assert iso[2:] == ["5 kVrms", "5.7 kVrms", "better"]


def test_compare_selected_parts_priority_1_and_2_only(client):
    a, b, c = ids(client, "SR8641", "SA7741", "SC141")
    html = client.get(f"/compare?pick={a}&pick={b}&pick={c}&baseline={a}").text
    assert "CMTI" in html and "Data rate" in html  # P1 / P2
    assert "Pulse width distortion" not in html  # P3 excluded
    data = client.get(f"/api/compare?ids={a},{b},{c}").json()
    assert {r["priority"] for r in data["rows"]} == {1, 2}


def test_compare_downloads(client):
    a, b = ids(client, "SR8233", "SA21520")
    x = client.get(f"/compare.xlsx?ids={a},{b}&baseline={a}")
    ws = load_workbook(io.BytesIO(x.content)).active
    assert any(c.value == "Propagation delay (typ, else max)" for row in ws.iter_rows() for c in row)
    text = client.get(f"/compare.csv?ids={a},{b}").text
    assert "SynAlpha SA21520" in text
    assert client.get("/compare.csv?ids=").status_code == 400


def test_api_crosses(client):
    (ref,) = ids(client, "SR8286")
    data = client.get(f"/api/parts/{ref}/crosses").json()
    assert set(data) == {"fff", "functional", "package"}
    assert len(data["fff"][0]["top_differences"]) <= 5


def test_reject_fff_and_undo(client):
    ref, cand = ids(client, "SR8641", "SC141")
    html = client.get(f"/parts/{ref}").text
    row = html.split(f'id="row-{cand}"')[1].split("</tr>")[0]
    assert "FFF" in row and "Reject as FFF" in html
    r = client.post(
        "/feedback", data={"reference_id": ref, "candidate_id": cand, "verdict": "reject", "note": "too slow"}
    )
    assert r.status_code == 200  # redirect followed back to the part page
    html = client.get(f"/parts/{ref}").text
    row = html.split(f'id="row-{cand}"')[1].split("</tr>")[0]
    assert "Package equivalent" in row and "FFF rejected in review" in row
    assert "too slow" in html and "Local user" in html
    data = client.get(f"/api/parts/{ref}/crosses").json()
    sc141 = next(x for x in data["package"] if x["candidate"]["label"].endswith("SC141"))
    assert sc141["fff_status"] == "rejected" and sc141["feedback"]["note"] == "too slow"
    client.post("/feedback", data={"reference_id": ref, "candidate_id": cand, "verdict": "clear"})
    data = client.get(f"/api/parts/{ref}/crosses").json()
    assert any(x["candidate"]["label"].endswith("SC141") for x in data["fff"])


def test_feedback_reject_moves_out_of_fff(client):
    ref, cand = ids(client, "SR8641", "SC141")
    client.post("/feedback", data={"reference_id": ref, "candidate_id": cand, "verdict": "reject", "note": "too slow"})
    data = client.get(f"/api/parts/{ref}/crosses").json()
    assert all(not x["candidate"]["label"].endswith("SC141") for x in data["fff"])
    assert "Rejected as FFF (1)" in client.get(f"/parts/{ref}").text
    for verdict in ("maybe", "accept"):
        r = client.post("/feedback", data={"reference_id": ref, "candidate_id": cand, "verdict": verdict})
        assert r.status_code == 400


def test_partial_part_numbers(client):
    from app.services.queries import normalize_pn

    assert normalize_pn(" si-86 41 ") == "SI8641"
    rows = client.get("/api/parts?q=8641").json()
    mpns = [p["mpn"] for p in rows]
    assert "SR8641" in mpns and "SD8641A" in mpns and "SB8641N" in mpns
    assert [p["mpn"] for p in client.get("/api/parts?q=sr86*41").json()][:1] == ["SR8641"]  # wildcard, prefix first
    assert {p["mpn"] for p in client.get("/api/parts?q=sa-77").json()} >= {"SA7741", "SA7740N", "SA7721"}


def test_search_page_covers_all_companies(client):
    html = client.get("/parts?q=7741").text
    assert "SA7741" in html  # competitor part found without changing the filter
    start = client.get("/parts").text  # one taxonomy: category tree, facets, capability badges
    assert 'class="cat-tree"' in start and "Data available" in start and "Isolation" in start
    fam = client.get("/parts?family=digital_isolator").text  # old family link still lands on the subfamily
    assert "SynRef" in fam and "Digital isolators" in fam
    one = client.get("/parts?family=digital_isolator&company=SynRef").text
    assert "SR8641" in one and "SA7741X" not in one and "Crosses" in one and "Reset all" in one


def test_suggestions_fragment(client):
    html = client.get("/suggest?q=864").text
    assert "SR8641" in html and "suggest-list" in html and "<html" not in html
    assert client.get("/suggest?q=8").text.strip() == ""  # at least two characters
    assert "No part numbers contain" in client.get("/suggest?q=zzzz").text


def test_result_table_filters_by_company(client):
    (ref,) = ids(client, "SR8641")
    html = client.get(f"/parts/{ref}").text
    assert 'class="company-filter"' in html and "SynAlpha (" in html
    table = (
        client.get(f"/parts/{ref}?company=SynBeta")
        .text.split('<table class="results">')[1]
        .split('<details class="category-view">')[0]
    )
    assert "SB8441" in table and "SA7741" not in table


def test_result_table_cross_types(client):
    ref, exact, better, info_worse, other_out, other_pins = ids(
        client, "SR8641", "SA7741X", "SD8641A", "SC141", "SB8441", "SB8641N"
    )

    def table(query):
        html = client.get(f"/parts/{ref}?all=true{query}").text
        return html, html.split('<table class="results">')[1].split('<details class="category-view">')[0]

    _, t = table("")  # non-automotive crosses of a non-automotive reference

    def row(pid, t=t):
        return t.split(f'id="row-{pid}"')[1].split("</tr>")[0]

    assert "Drop-in" in row(exact) and "spec-exact" in row(exact) and "p2p yes" in row(exact)
    assert "Drop-in" in row(info_worse) and "spec-exact" in row(info_worse)  # only priority 2 worse
    assert "Package eq." in row(other_out) and "p2p yes" in row(other_out)
    assert "Functional" in row(other_pins) and "p2p no" in row(other_pins)
    assert t.index(f'id="row-{info_worse}"') < t.index(f'id="row-{other_pins}"') < t.index(f'id="row-{other_out}"')
    _, auto = table("&grade=automotive")
    assert "Drop-in" in row(better, auto) and "spec-over" in row(better, auto)
    page = client.get(f"/parts/{ref}?type=functional").text
    only = page.split('<table class="results">')[1].split('<details class="category-view">')[0]
    assert f'id="row-{other_pins}"' in only and f'id="row-{other_out}"' not in only and f'id="row-{exact}"' not in only
    assert "Functional (6)" in page and "Drop-in (3)" in page  # SD140N: creepage is priority 3 now
    assert "Exception (listed, not a knock-out)" in client.get(f"/parts/{ref}?all=true").text


def test_deviation_setting(client):
    ref, cand = ids(client, "SR8640N", "SB8441")  # other package, one priority 1 deviation
    assert f'id="row-{cand}"' in client.get(f"/parts/{ref}?all=true").text
    assert f'id="row-{cand}"' not in client.get(f"/parts/{ref}?dev=0&all=true").text
    assert "SB8441" not in client.get(f"/parts/{ref}/replacements.csv?dev=0").text


def test_competitor_part_finds_reference_replacements(client):
    (comp,) = ids(client, "SA7741X")  # identical to SR8641, but obsolete
    html = client.get(f"/parts/{comp}").text
    assert "Suitable Skyworks parts" in html and "Include other competitors" in html
    data = client.get(f"/api/parts/{comp}/crosses").json()
    fff = [x["candidate"]["label"] for x in data["fff"]]
    assert fff == ["SynRef SR8641"]
    every = {x["candidate"]["label"] for cat in data.values() for x in cat}
    assert all(label.startswith("SynRef") for label in every)  # only reference parts by default
    wide = client.get(f"/parts/{comp}?scope=all").text
    assert "Replacements" in wide and "SB8441" in wide


def test_reference_part_page_unchanged(client):
    (ref,) = ids(client, "SR8641")
    html = client.get(f"/parts/{ref}").text
    assert "<h2>Replacements</h2>" in html and "Suitable Skyworks parts" not in html


def test_suggestions_mark_competitor_parts(client):
    html = client.get("/suggest?q=7741").text
    assert "find a Skyworks part" in html


def test_finder_start_page(client):
    html = client.get("/").text
    assert "Search &amp; replacements" in html and 'action="/find"' in html
    # three starting actions and one definition of the cross types
    assert all(x in html for x in ('href="/compare"', 'href="/blocks"', 'href="/blocks/cross-sell"'))
    assert "every priority 1 parameter equal or better" in html


def test_find_full_part_number_opens_its_page(client):
    (comp,) = ids(client, "SA7741X")
    r = client.get("/find?q=sa-7741x", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == f"/parts/{comp}"


def test_find_partial_lists_competitors_first(client):
    html = client.get("/find?q=7741").text
    assert "Tap a part to see suitable Skyworks parts" in html and "SA7741X" in html
    assert "No part numbers contain" in client.get("/find?q=ADUM1401").text
    assert client.get("/find?q=", follow_redirects=False).status_code == 303


def test_result_table_download(client):
    (ref,) = ids(client, "SR8641")
    r = client.get(f"/parts/{ref}/replacements.xlsx")
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    header = next(i for i, row in enumerate(rows) if row[0] == "Part")
    verdicts = {row[0]: row[2] for row in rows[header + 1 :]}
    assert verdicts["SA7741"] == "Drop-in (FFF)" and verdicts["SB8441"].startswith("Package equivalent")
    assert verdicts["SB8641N"].startswith("Functional equivalent")
    text = client.get(f"/parts/{ref}/replacements.csv?type=fff").text
    assert "SA7741" in text and "SB8641N" not in text


def test_cmti_explorer_page(client):
    ref, cand = ids(client, "SR8641", "SA7741")
    html = client.get(f"/engineering/cmti_kv_us?ids={ref},{cand}").text
    assert "CMTI" in html and "Datasheet guaranteed" in html and "Engineering estimate" in html
    devices = json.loads(html.split("data-devices='")[1].split("'")[0])
    assert [d["mpn"] for d in devices] == ["SR8641", "SA7741"]  # reference first
    assert all(isinstance(d["values"]["cmti_kv_us"], int | float) for d in devices)
    assert f'href="/engineering/cmti_kv_us?ids={ref},' in client.get(f"/parts/{ref}").text
    r = client.get(f"/impact?ids={ref},{cand}", follow_redirects=False)  # old address
    assert r.headers["location"] == f"/engineering/cmti_kv_us?ids={ref},{cand}"


def test_engineering_notes(client):
    html = client.get("/engineering").text
    assert "Specification guides" in html and "/engineering/creepage_mm" in html
    for key in ("isolation_rating_kvrms", "viorm_vpk", "prop_delay_ns", "peak_sink_a", "enable_pins"):
        assert client.get(f"/engineering/{key}").status_code == 200, key
    assert client.get("/engineering/no_such_thing").status_code == 404


def test_deviation_links_open_the_note_with_both_parts(client):
    ref, cand = ids(client, "SR8640N", "SB8441")  # default output deviates
    page = client.get(f"/parts/{ref}?all=true").text
    # lower-ranked rows load their details when opened: the page links the fragment, the fragment has the links
    detail = page if f"cand={cand}" in page else client.get(f"/parts/{ref}/detail/{cand}?all=true").text
    link = f"/engineering/default_output?ref={ref}&amp;cand={cand}"
    assert link in detail
    note = client.get(link.replace("&amp;", "&")).text
    assert "Your comparison" in note and "SR8640N" in note and "SB8441" in note and "different" in note


def test_grade_badges_and_filter(client):
    (ref,) = ids(client, "SR8641")
    html = client.get(f"/parts/{ref}").text
    assert "Non-automotive reference" in html and "All grades" not in html
    assert "SD8641A" not in html.split('<details class="category-view">')[0]  # non-automotive by default
    auto = client.get(f"/parts/{ref}?grade=automotive").text
    assert 'class="badge grade-auto"' in auto
    assert "SD8641A" in auto and "SC141" not in auto.split('<details class="category-view">')[0]
    ind = client.get(f"/parts/{ref}?grade=industrial").text
    assert "SD8641A" not in ind.split('<details class="category-view">')[0]


def test_automotive_reference_banner(client):
    (ref,) = ids(client, "SD8641A")
    html = client.get(f"/parts/{ref}").text
    assert "Automotive reference" in html and "Filter by grade" not in html


def test_navigation_sections_and_old_urls(client):
    """Sections in the top row, the section's pages in the second: every old URL lands in the right section and item."""
    for url, section, item in [("/", "Products", "Search &amp; replacements"),
                               ("/parts", "Products", "Search &amp; replacements"),
                               ("/compare", "Products", "Compare"),
                               ("/broadcast", "Systems &amp; applications", "Application library"),
                               ("/broadcast/configurations", "Systems &amp; applications", "Application library"),
                               ("/blocks", "Systems &amp; applications", "Application library"),
                               ("/blocks/coverage", "Systems &amp; applications", "Application library"),
                               ("/business", "Opportunities", "Gaps &amp; new technology"),
                               ("/blocks/cross-sell", "Systems &amp; applications", "Cross-sell"),
                               ("/business/attach", "Systems &amp; applications", "Cross-sell"),
                               ("/portfolio", "Systems &amp; applications", "Cross-sell"),
                               ("/business/rf-coverage", "Systems &amp; applications", "RF coverage"),
                               ("/business/automotive", "Systems &amp; applications", "Automotive map"),
                               ("/business/tam", "Future growth", "TAM estimates"),
                               ("/manage/overview", "Data management", "Overview"),
                               ("/manage/verify", "Data management", "Site data"),
                               ("/manage/feedback", "Data management", "Site data"),
                               ("/manage/guest-input", "Data management", "Site data"),
                               ("/business/competitors", "Opportunities", "Competitive analytics"),
                               ("/business/markets", "Future growth", "Markets &amp; customers"),
                               ("/business/reports", "Future growth", "Research &amp; ideas"),
                               ("/library", "Systems &amp; applications", "Systems library"),
                               ("/library/v2x_frontend", "Systems &amp; applications", "Systems library"),
                               ("/visuals", "Engineering knowledge", "Interactive visuals"),
                               ("/rf-bands", "Engineering knowledge", "Interactive visuals"),
                               ("/knowledge", "Engineering knowledge", "General knowledge"),
                               ("/rf-notes", "Engineering knowledge", "General knowledge"),
                               ("/technical", "Engineering knowledge", "General knowledge"),
                               ("/engineering", "Engineering knowledge", "General knowledge"),
                               ("/engineering/cmti_kv_us", "Engineering knowledge", "General knowledge"),
                               ("/business/capabilities", "Engineering knowledge", "General knowledge"),
                               ("/rules", "Data management", "Methodology"),
                               ("/requests", "Data management", "Site data")]:  # fmt: skip
        r = client.get(url)
        assert r.status_code == 200, url
        assert f'aria-current="page" class="on">{item}<' in r.text, url
        assert re.search(r'aria-current="true"[^>]*><svg[^>]*>.*?</svg>' + re.escape(section) + "<", r.text), url
        assert 'class="sub-nav"' in r.text and 'id="sidebar"' not in r.text, url
    # in-page tabs connect related views instead of a long row of sibling tabs
    html = client.get("/blocks/coverage").text
    assert 'href="/business/applications"' in html and 'aria-current="page" class="on">Technical coverage<' in html
    assert (
        'id="global-search"' in client.get("/broadcast").text and 'id="global-search"' not in client.get("/parts").text
    )
    # Engineering knowledge (the Systems library moved to Systems & applications, owner 23 Oct 2026; the QSPICE
    # library added 27 Oct 2026)
    from app.web import nav

    knowledge = next(s for s in nav.NAV if s[0] == "knowledge")
    assert [i[0] for i in knowledge[2]] == ["Interactive visuals", "General knowledge", "QSPICE library"]
    systems = next(s for s in nav.NAV if s[0] == "systems")
    assert [i[0] for i in systems[2]] == ["Systems library", "Application library", "Application workspace",
                                         "Cross-sell", "RF coverage", "Automotive map"]
    rf = client.get("/rf-notes").text
    assert 'aria-current="page" class="on">RF fundamentals<' in rf and 'href="/engineering"' in rf
    assert 'aria-current="page" class="on">Specification guides<' in client.get("/engineering/cmti_kv_us").text
    bands = client.get("/rf-bands").text
    assert 'aria-current="page" class="on">RF bands<' in bands and 'href="/visuals"' in bands
    assert nav.page_title("/rf-notes").startswith("RF fundamentals – Engineering knowledge")
    assert "Technical library" not in client.get("/library").text


def test_best_and_next_best_options_are_flagged(client):
    (ref,) = ids(client, "SR8641")
    html = client.get(f"/parts/{ref}").text.split('<details class="category-view">')[0]
    # best and next best per cross type, nothing else until "show all"
    assert html.count("Best drop-in</span>") == 1 and html.count("Next best drop-in</span>") == 1
    assert html.count('class="result-row') <= 6 and "Show all results (" in html and "Other options (" not in html
    full = client.get(f"/parts/{ref}?all=true").text.split('<details class="category-view">')[0]
    assert "Other options (" in full and full.count('class="result-row') > html.count('class="result-row')


def test_compare_page_layout_and_typed_parts(client):
    html = client.get("/compare?add=SR8641,SA7741").text
    assert html.index('name="add"') < html.index("compare-table")
    assert "Pick from the list" not in html  # the series list was removed; search and part pages feed Compare
    assert "Not found" in client.get("/compare?add=NOSUCHPART").text


def test_series_and_sister_family_prefixes():
    from app.services.alternates import sister_prefix
    from app.web.routes import series_of

    assert series_of("SI86S641BE-IS2") == "SI86S64x" and series_of("NSI8241W1-DSWR") == "NSI824x"
    assert sister_prefix("SI8641BD-IS")[0] == "SI86S641" and sister_prefix("SI86S641BE-IS2")[0] == "SI8641"
    assert sister_prefix("SI8234AB-AM") is None


def test_part_page_loads_heavy_sections_on_demand(client):
    ref, cand = ids(client, "SR8640N", "SB8441")
    page = client.get(f"/parts/{ref}?all=true").text
    assert f'hx-get="/parts/{ref}/grouped' in page and "data-src=" in page  # loaded when opened
    grouped = client.get(f"/parts/{ref}/grouped").text
    assert "Top 5 differences" in grouped
    assert client.get(f"/parts/{ref}/detail/999999").status_code == 404


def test_search_covers_the_manufacturer_catalogues(client):
    """Catalogue-only parts (no parameters or crosses yet) are found by the part-number search and suggestions."""
    html = client.get("/find?q=SKY5A3200").text
    assert "Skyworks &amp; Qorvo catalogue" in html and "<b>SKY5A3200</b>" in html and "Front-end Modules" in html
    assert "SKY5A3200" in client.get("/suggest?q=sky5a32").text and "sg-cat" in client.get("/suggest?q=sky5a32").text
    assert "in the Skyworks and Qorvo catalogues" in client.get("/find?q=zzzzqqq").text


def test_cross_row_links_part_and_has_details_toggle(client):
    (ref,) = ids(client, "SR8641")
    html = client.get(f"/parts/{ref}").text
    assert 'class="cand-link"' in html and 'class="detail-toggle"' in html
    assert "Electrical: " in html  # electrical verdict kept apart from package and pinout


def test_footer_and_browse_off_the_menu(client):
    """Owner, 24 Oct 2026: Browse products is off the menu (too long); the page still opens from links."""
    html = client.get("/parts").text
    assert 'class="site-foot"' in html and "Cross rules" in html and "<h1>Browse products</h1>" in html
    assert 'href="/parts" title="Every part by function' not in html
    assert 'aria-current="page" class="on">Search &amp; replacements' in html


def test_suggestions_are_a_combobox(client):
    assert 'role="combobox"' in client.get("/").text
    assert 'id="sg-p' in client.get("/suggest?q=7741").text


def test_general_knowledge_page(client):
    """General knowledge holds RF fundamentals as a sub-category, the specification guides, technology platforms and
    the search, and lists the systems that use them."""
    html = client.get("/knowledge").text
    assert "<h1>General knowledge</h1>" in html
    for href in ('href="/rf-notes"', 'href="/rf-notes#cat-rf_specs"', 'href="/engineering"',
                 'href="/engineering/cmti_kv_us"', 'href="/business/capabilities"',
                 'action="/technical"', 'href="/library/v2x_frontend"'):  # fmt: skip
        assert href in html, href


def test_visuals_include_rf_bands():
    from pathlib import Path

    js = (Path(__file__).resolve().parents[1] / "app" / "web" / "static" / "visuals.js").read_text()
    assert 'href="/rf-bands"' in js  # the RF band chart is listed with the visuals


def test_systems_library_links_to_knowledge_and_visuals(client):
    from app.services import engineering, library, rf_notes

    ids = {t["id"] for t in rf_notes.topics()}
    assert set(library.RF_NOTES) <= ids and set(library.SPEC_GUIDES) <= set(engineering.NOTES)
    static = library.PROJECT_ROOT / "app" / "web" / "static"  # visuals.js plus visuals in their own files (vz/)
    js = "".join(f.read_text() for f in [static / "visuals.js", *sorted(static.glob("vz/*.js"))])
    assert all(f'id: "{k}"' in js for k in library.VISUALS if k != "rf-bands")
    v2x = library.related("v2x_frontend")
    hrefs = {x["href"] for x in v2x["knowledge"] + v2x["visuals"]}
    assert "/rf-notes#acoustic_filters_duplexers" in hrefs and "/visuals#filter" in hrefs and "/rf-bands" in hrefs
    traction = library.related("traction_asil_d")
    assert "/engineering/cmti_kv_us" in {x["href"] for x in traction["knowledge"]}
    assert "/visuals#cmti" in {x["href"] for x in traction["visuals"]}
    assert not any(x["href"].startswith("/rf-notes") for x in traction["knowledge"][:5])  # isolation, not RF
    page = client.get("/library/v2x_frontend").text
    assert 'class="kn-links' in page and 'href="/rf-notes#acoustic_filters_duplexers"' in page
    assert 'href="/visuals#filter"' in page and ">basics</a>" in page


def test_rf_by_band_table_has_a_key_and_marks_each_line(client):
    html = client.get("/portfolio").text
    table = html.split('<h2 id="rf">')[1].split('<h2 id="overlap">')[0]
    assert 'class="pf-bkey' in table and "Skyworks parts" in table and "a gap for us" in table
    assert 'class="pf-n pf-sky' in table and 'class="pf-n pf-qor' in table and not re.search(r"\d / \d", table)


# ---- from tests/test_fb_1022.py --------------------------------------------------------------------------------------
# Owner comments of 22 Oct 2026 (daily update): precision fixes, layout changes and new sections.

ROOT = Path(__file__).resolve().parents[1]


def test_sso8_creepage_is_at_least_not_measured():
    files = [
        ROOT / "data/business/cases/hv-extra-wide-isolation.yaml",
        ROOT / "app/web/templates/business/reports/iso7821dww-isolation.html",
        ROOT / "app/web/templates/business/reports/hv-extra-wide-isolation.background.html",
        ROOT / "data/business/reports.yaml",
    ]
    for f in files:
        text = f.read_text()
        assert "measured 8.5" not in text and "measures 8.5" not in text and "8.5 mm measured" not in text, f
    bg = files[2].read_text()
    assert "7. Summary and Future Outlook" in bg
    assert "Skyworks is currently developing a new ultra-wide-body (UWB) SOIC package" in bg


def test_walk_market_size_points_to_tam_and_is_not_logged(client, session, real):
    from sqlalchemy import select

    from app.models import SiteFeedback
    from app.services import walk

    comp = find_id(client, "ISO7741DBQR")
    html = client.get(f"/parts/{comp}/walk?upto=5").text
    assert "ASP and application volumes" not in html
    assert "no public market-size estimate yet" in html or "public estimate" in html
    texts = session.scalars(select(SiteFeedback.text).where(SiteFeedback.author == walk.AUTHOR)).all()
    assert not any("market size" in t or "market-size" in t for t in texts)


def test_no_feedback_sections_at_the_bottom(client):
    caps = client.get("/business/capabilities").text
    assert 'action="/business/capabilities/feedback"' not in caps and 'id="feedback"' not in caps
    rep = client.get("/business/reports/filters-1-11ghz").text
    assert "/ask" not in rep and 'class="fb-open"' in rep  # only the floating button
    # old form targets still answer
    r = client.post("/business/capabilities/feedback", data={"kind": "idea", "text": "x"}, follow_redirects=False)
    assert r.status_code == 303


def test_walk_cards_short_heading_full_text_and_next_on_top(client, real):
    from app.services import walk

    tws = "True-wireless (TWS) earbuds with charging case"
    assert walk.short_name(tws) == "True-wireless earbuds with charging case"
    assert walk.short_name("AC solid-state circuit breaker (SSCB) with air-gap relay, AC/DC current sensing") == (
        "AC solid-state circuit breaker"
    )
    comp = find_id(client, "ISO7741DBQR")
    html = client.get(f"/parts/{comp}/walk").text
    assert html.index("Next: 3 Best of ours around it") < html.index('id="step-1"')  # with the step buttons on top
    assert html.count("Next: 3") == 1
    assert 'class="link-btn walk-name"' in html and 'id="walk-full"' in html
    assert "more</li>" not in html  # nothing cut off
    css = (ROOT / "app/web/static/app.css").read_text()
    assert "text-overflow: ellipsis; white-space: nowrap; }\n.walk-also" not in css


def test_band_tables_do_not_scroll(client):
    rf = client.get("/business/rf-coverage").text
    assert '<div class="fit-wrap"><table class="rf-matrix stack-table rf-small">' in rf and 'data-label="' in rf
    pf = client.get("/portfolio").text
    assert 'class="pf-band small stack-table"' in pf
    css = (ROOT / "app/web/static/app.css").read_text()
    assert ".stack-table thead { display: none; }" in css


def test_agree_disagree_shown_without_react_click(client):
    html = client.get("/library/v2x_frontend").text
    assert 'class="pc-react pc-inline"' in html and ">Agree</button>" in html and ">Disagree</button>" in html
    assert '<details class="pc-react"' not in html and ">react<" not in html
    for path in ("app/web/templates/business/_prelim.html", "app/web/templates/reviews/detail.html"):
        src = (ROOT / path).read_text()
        assert '<details class="pc-react"' not in src and "pc-inline" in src
    r = client.post("/library/v2x_frontend/react", data={"item": "x", "verdict": "agree"}, follow_redirects=False)
    assert r.status_code in (303, 400, 404)


def test_markets_filters_compact_with_more_filters(client):
    html = client.get("/business/markets").text
    assert 'class="more-filters"><summary class="small">More filters' in html
    head, more = html.split('class="more-filters"', 1)
    assert ">Leap</span>" in head and "Qualify for automotive" in more.split("</details>")[0]
    assert 'class="more-filters" open' in client.get("/business/markets?view=multi").text


def test_feedback_window_has_dictation(client):
    html = client.get("/").text
    assert '<p class="fb-dict small" hidden>' in html and "fb-mic" in html  # shown by the script only if supported
    js = (ROOT / "app/web/static/feedback.js").read_text()
    assert "window.SpeechRecognition || window.webkitSpeechRecognition" in js and "if (!SR || !root) return;" in js


def test_research_lists_product_ideas_as_feasibility_studies(client):
    html = client.get("/business/reports").text
    assert "Product idea" in html and "Feasibility study" in html
    assert ">Report</span>" not in html and ">Type</span>" not in html  # no report / idea split
    assert ">Application</a>" in html and ">Feasibility</a>" in html
    assert "Reports <small" not in client.get("/business/reports?view=feasibility").text
    page = client.get("/business/reports/filters-1-11ghz").text
    assert "Feasibility study" in page and ">Report</span>" not in page
    r = client.post("/business/schedule", data={"title": "Filters", "source": "research:filters-1-11ghz"},
                    follow_redirects=False)  # fmt: skip
    assert r.status_code == 303
    assert "Business case review (scheduled)" in client.get("/business/reports/filters-1-11ghz").text


def test_business_case_asks_feasibility_price_manufacturing_markets_competition(client):
    import re

    from app.services import cases

    keys = [k for k, *_ in cases.FIELDS]
    for k in ("feasibility", "price", "manufacturing", "inhouse", "markets", "market_size", "competitors",
              "competitor_prices", "competitor_share", "anything_else"):  # fmt: skip
        assert k in keys, k
    page = client.get("/business/cases/new?fn=pol_ldo").text
    assert "Is the technology feasible?" in page and "Anything else to add?" in page
    assert 'data-field="market_size"' in page and "Manufacturing capability" in page
    r = client.post(
        "/business/cases",
        data={
            "title": "LDO case",
            "fn": "pol_ldo",
            "feasibility": "Existing die",
            "manufacturing": "Open: not public",
            "anything_else": "Add export",
        },
    )
    assert r.status_code in (200, 303)
    cid = int(re.search(r"/business/cases/(\d+)", client.get("/business/cases").text).group(1))
    brief = client.get(f"/business/cases/{cid}/brief.md").text
    assert "Existing die" in brief and "Manufacturing" in brief and "Anything else to add?" in brief
    assert "public prices" in brief and "share of the" in brief
    sg = re.search(r'<script type="application/json" id="sg-data">(.*?)</script>', page, re.S)
    assert sg and "market_size" in sg.group(1)  # a hint for market size (estimate or "add one on TAM estimates")


def test_price_estimate_variants_sop_and_break_rows(client, session, real, part):
    from datetime import date

    html = client.get("/pricing/estimate", params={"name": "ISO7741"}).text
    assert 'class="pe-variants"' in html and "ISO7741DBQR" in html and "ISO7741DWR" in html
    assert 'class="outline small-btn pe-add"' in html and ">or</span>" in html and 'name="bq"' in html
    p = part("ISO7741DWR")
    p.release_date = date(2016, 3, 1)
    session.commit()
    html = client.get("/pricing/estimate", params={"name": "ISO7741DWR"}).text
    assert 'name="sop" type="number" min="1980" max="2050" value="2016"' in html
    assert "year of the datasheet&#39;s first revision" in html or "year of the datasheet's first revision" in html
    # rows and pasted text together make the breaks
    q = [("family", "digital_isolator"), ("name", "X1"), ("bq", "1"), ("bp", "3.10"), ("bq", "1000"), ("bp", "1.90"),
         ("breaks", "100 2.50")]  # fmt: skip
    html = client.get("/pricing/estimate", params=q).text
    assert 'id="price-model"' in html and 'value="1000"' in html and "100 2.50" in html


def test_library_focus_areas_and_area_requests(client, session):
    from datetime import date

    from app.services import blocks, inbox

    assert blocks.focus_of("Quadcopter drone flight controller", "", "Consumer") == ["drones"]
    assert blocks.focus_of("5G smartphone modem-to-antenna", "", "Consumer") == ["cellular"]
    assert "automotive" in blocks.focus_of("Body controller", "", "Automotive")
    assert blocks.focus_of("Cell site gateway router", "Communications equipment", "Industrial") == ["cellular"]
    assert blocks.focus_of("Smart thermostat", "Appliances & building", "Consumer") == []
    html = client.get("/blocks").text
    order = [html.index(f'id="area-{k}"') for k in ("automotive", "datacenter", "drones", "cellular")]
    assert order == sorted(order) and order[-1] < html.index('id="other-areas"')
    assert '<details class="lib-other" id="other-areas">' in html and 'action="/blocks/area-request"' in html
    r = client.post("/blocks/area-request", data={"area": "Robotics", "why": "Humanoid robots"}, follow_redirects=False)
    assert r.status_code == 303
    assert client.post("/blocks/area-request", data={"area": " "}).status_code == 400
    assert "Robotics" in client.get("/manage/reviews?type=area").text
    page = client.get("/blocks/area-requests").text
    assert "Robotics" in page and "Humanoid robots" in page
    text = inbox.collect(session, date.today())
    assert "## Application area requests" in text and "Robotics" in text and "why: Humanoid robots" in text
    rid = int(page.split('action="/blocks/area-requests/')[1].split('"')[0])
    r = client.post(f"/blocks/area-requests/{rid}", data={"status": "planned"}, follow_redirects=False)
    assert r.status_code == 303
    assert client.post(f"/blocks/area-requests/{rid}", data={"status": "x"}).status_code == 400
    daily = (ROOT / "docs/daily_update.md").read_text()
    assert ("Deep dives for data centres and drones (main areas outside automotive and cellular): a few documents per "
            "run") in daily  # fmt: skip
